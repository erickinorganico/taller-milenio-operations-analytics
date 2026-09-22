"""Local Excel/CSV boundary for small client operations extracts.

This module is deliberately an input boundary, not a CRM or an accounting
connector.  It accepts one workbook or a directory of CSV files, validates a
small, stable contract, and returns JSON-friendly canonical tables.  Missing
optional values stay ``None`` so downstream analytics can report ``unknown``
instead of manufacturing zeroes.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import openpyxl
import xlsxwriter


SCHEMA_VERSION = 1
UTC = timezone.utc
TABLES = ("orders", "invoices", "payments", "inventory")
SHEET_NAMES = {
    "orders": "Ordenes",
    "invoices": "Facturas",
    "payments": "Pagos",
    "inventory": "Inventario",
}
CONFIG_SHEET = "Config"
INSTRUCTIONS_SHEET = "Instrucciones"

TABLE_SPECS: dict[str, dict[str, Any]] = {
    "orders": {
        "required": ("order_id", "customer_ref", "segment", "received_at", "status"),
        "fields": (
            "order_id",
            "customer_ref",
            "segment",
            "received_at",
            "delivered_at",
            "status",
            "promised_at",
            "block_reason",
            "technician_ref",
            "estimated_hours",
        ),
        "headers": (
            "ID de orden",
            "Referencia cliente",
            "Segmento",
            "Recibida en",
            "Entregada en",
            "Estado",
            "Prometida para",
            "Motivo de bloqueo",
            "Referencia técnico",
            "Horas estimadas",
        ),
        "aliases": {
            "order_id": ("order_id", "orden_id", "id_orden", "id de orden", "folio_orden"),
            "customer_ref": (
                "customer_ref", "cliente_ref", "customer_id", "cliente_id",
                "referencia_cliente", "referencia cliente", "id_cliente",
            ),
            "segment": ("segment", "segmento", "tipo_cliente"),
            "received_at": ("received_at", "received", "recibida_en", "fecha_recepcion", "fecha_recibida"),
            "delivered_at": ("delivered_at", "delivered", "entregada_en", "fecha_entrega", "fecha_entregada"),
            "status": ("status", "estado"),
            "promised_at": ("promised_at", "promised", "prometida_para", "fecha_prometida"),
            "block_reason": ("block_reason", "motivo_bloqueo", "motivo_de_bloqueo", "bloqueo"),
            "technician_ref": ("technician_ref", "technician_id", "tecnico_ref", "referencia_tecnico", "tecnico"),
            "estimated_hours": ("estimated_hours", "horas_estimadas", "horas_estimadas_trabajo"),
        },
        "enums": {"segment": ("particular", "fleet"), "status": ("received", "in_service", "waiting_parts", "rework", "ready", "delivered", "cancelled")},
    },
    "invoices": {
        "required": ("invoice_id", "customer_ref", "issued_at", "amount_mxn", "status"),
        "fields": ("invoice_id", "order_id", "customer_ref", "issued_at", "due_at", "amount_mxn", "status"),
        "headers": ("ID de factura", "ID de orden", "Referencia cliente", "Emitida en", "Vencimiento", "Importe MXN", "Estado"),
        "aliases": {
            "invoice_id": ("invoice_id", "factura_id", "id_factura", "id de factura", "folio_factura"),
            "order_id": ("order_id", "orden_id", "id_orden", "id de orden", "folio_orden"),
            "customer_ref": ("customer_ref", "cliente_ref", "customer_id", "cliente_id", "referencia_cliente", "referencia cliente", "id_cliente"),
            "issued_at": ("issued_at", "issued", "emitida_en", "fecha_emision", "fecha_emitida"),
            "due_at": ("due_at", "due", "vencimiento", "fecha_vencimiento", "vence_en"),
            "amount_mxn": ("amount_mxn", "amount", "importe_mxn", "importe", "monto_mxn", "monto"),
            "status": ("status", "estado"),
        },
        "enums": {"status": ("issued", "void")},
    },
    "payments": {
        "required": ("payment_id", "invoice_id", "paid_at", "amount_mxn", "method"),
        "fields": ("payment_id", "invoice_id", "paid_at", "amount_mxn", "method"),
        "headers": ("ID de pago", "ID de factura", "Pagado en", "Importe MXN", "Método"),
        "aliases": {
            "payment_id": ("payment_id", "pago_id", "id_pago", "id de pago", "folio_pago"),
            "invoice_id": ("invoice_id", "factura_id", "id_factura", "id de factura", "folio_factura"),
            "paid_at": ("paid_at", "paid", "pagado_en", "fecha_pago", "fecha_pagado"),
            "amount_mxn": ("amount_mxn", "amount", "importe_mxn", "importe", "monto_mxn", "monto"),
            "method": ("method", "metodo", "método", "forma_pago", "forma de pago"),
        },
        "enums": {"method": ("cash", "card", "transfer")},
    },
    "inventory": {
        "required": ("part_id", "description", "on_hand", "reserved", "reorder_point"),
        "fields": ("part_id", "description", "on_hand", "reserved", "reorder_point", "unit_cost_mxn"),
        "headers": ("ID de parte", "Descripción", "Existencia", "Reservado", "Punto de reorden", "Costo unitario MXN"),
        "aliases": {
            "part_id": ("part_id", "parte_id", "id_parte", "id de parte", "sku", "codigo_parte"),
            "description": ("description", "descripcion", "descripción", "nombre_parte", "parte"),
            "on_hand": ("on_hand", "existencia", "disponible", "cantidad_existencia", "stock"),
            "reserved": ("reserved", "reservado", "cantidad_reservada"),
            "reorder_point": ("reorder_point", "punto_reorden", "punto de reorden", "minimo", "mínimo"),
            "unit_cost_mxn": ("unit_cost_mxn", "unit_cost", "costo_unitario_mxn", "costo_unitario", "costo"),
        },
        "enums": {},
    },
}

CONFIG_ALIASES = {
    "business_name": ("business_name", "nombre_negocio", "nombre_empresa", "negocio"),
    "as_of": ("as_of", "corte", "fecha_corte", "fecha_as_of", "corte_en"),
    "snapshot_id": ("snapshot_id", "id_snapshot", "id_corte", "snapshot", "identificador_corte"),
    "synthetic": ("synthetic", "sintetico", "sintética", "sintetica", "datos_sinteticos", "datos_sintéticos"),
}
CONFIG_FIELDS = ("business_name", "as_of", "snapshot_id", "synthetic")
IGNORED_SHEETS = {INSTRUCTIONS_SHEET.lower(), "readme", "instructions"}


def _norm(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.strip().casefold()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: Any) -> str | None:
    if _blank(value):
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _iso(value: datetime) -> str:
    value = value.astimezone(UTC).replace(microsecond=0)
    return value.isoformat().replace("+00:00", "Z")


def _issue(sheet: str, row: int | None, column: str | int | None, code: str, message: str) -> dict[str, Any]:
    return {"sheet": sheet, "row": row, "column": column, "code": code, "message": message}


class ClientInputError(ValueError):
    """Validation failure with row/cell-addressable issues."""

    def __init__(self, issues: Sequence[Mapping[str, Any]], message: str = "La entrada del cliente no cumple el contrato"):
        self.issues = [dict(item) for item in issues]
        detail = "; ".join(
            f"{item.get('sheet')}!{item.get('column') or '?'}{item.get('row') or ''}: {item.get('message')}"
            for item in self.issues[:8]
        )
        super().__init__(message if not detail else f"{message}: {detail}")


def _as_explicit_as_of(value: Any, *, sheet: str = "Config", row: int | None = None, column: str | int | None = "as_of", allow_naive_excel: bool = True) -> tuple[datetime | None, list[dict[str, Any]], bool]:
    """Parse the cutoff; Excel-native naive values are documented as UTC."""
    if _blank(value):
        return None, [_issue(sheet, row, column, "missing_as_of", "Config.as_of es obligatorio y debe incluir zona horaria")], False
    try:
        if isinstance(value, datetime):
            if value.tzinfo is None:
                if allow_naive_excel:
                    return value.replace(tzinfo=UTC), [], True
                return None, [_issue(sheet, row, column, "timezone_required", "La fecha debe incluir zona horaria")], False
            return value.astimezone(UTC), [], False
        if isinstance(value, date):
            if allow_naive_excel:
                return datetime.combine(value, time.min, tzinfo=UTC), [], True
            return None, [_issue(sheet, row, column, "timezone_required", "La fecha debe incluir zona horaria")], False
        if isinstance(value, str):
            parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00").replace("z", "+00:00"))
            if parsed.tzinfo is None:
                return None, [_issue(sheet, row, column, "timezone_required", "La fecha debe ser ISO-8601 con zona horaria")], False
            return parsed.astimezone(UTC), [], False
    except (TypeError, ValueError, OverflowError):
        pass
    return None, [_issue(sheet, row, column, "invalid_date", "Fecha inválida; use ISO-8601 con zona horaria")], False


def _parse_date(value: Any, *, sheet: str, row: int, column: str, field: str) -> tuple[str | None, list[dict[str, Any]], bool]:
    if _blank(value):
        return None, [], False
    parsed, issues, assumed = _as_explicit_as_of(value, sheet=sheet, row=row, column=column)
    if issues or parsed is None:
        if issues:
            return None, [_issue(sheet, row, column, issues[0]["code"], f"{field}: {issues[0]['message']}")], assumed
        return None, [_issue(sheet, row, column, "invalid_date", f"{field}: fecha inválida")], assumed
    return _iso(parsed), [], assumed


def _decimal(value: Any, *, sheet: str, row: int, column: str, field: str, nonnegative: bool = True, max_value: Decimal | None = None) -> tuple[Decimal | None, list[dict[str, Any]]]:
    if _blank(value):
        return None, []
    if isinstance(value, bool):
        return None, [_issue(sheet, row, column, "invalid_number", f"{field}: se esperaba un número")]
    if isinstance(value, float) and not math.isfinite(value):
        return None, [_issue(sheet, row, column, "invalid_money", f"{field}: NaN o infinito no está permitido")]
    try:
        text = str(value).strip()
    except (ValueError, OverflowError):
        return None, [_issue(sheet, row, column, "invalid_number", f"{field}: número demasiado grande")]
    text = text.replace("$", "").replace("MXN", "").strip()
    # Accept only correctly grouped thousands separators.  In particular,
    # ``1,2.34`` must never become ``12.34`` by silently removing commas.
    if "," in text:
        if not re.fullmatch(r"-?(?:[0-9]{1,3}(?:,[0-9]{3})+)(?:\.[0-9]+)?", text):
            return None, [_issue(sheet, row, column, "invalid_money", f"{field}: separadores de miles inválidos")]
        text = text.replace(",", "")
    if len(text) > 128:
        return None, [_issue(sheet, row, column, "invalid_number", f"{field}: número demasiado grande")]
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        return None, [_issue(sheet, row, column, "invalid_money", f"{field}: importe MXN inválido")]
    if not number.is_finite() or (nonnegative and number < 0):
        return None, [_issue(sheet, row, column, "invalid_money", f"{field}: importe MXN inválido o negativo")]
    if number.as_tuple().exponent < -128:
        return None, [_issue(sheet, row, column, "number_range", f"{field}: precisión fuera de rango")]
    if max_value is not None and number > max_value:
        return None, [_issue(sheet, row, column, "number_range", f"{field}: número fuera de rango")]
    return number, []


def _money_cents(value: Any, *, sheet: str, row: int, column: str, field: str, optional: bool = False) -> tuple[int | None, list[dict[str, Any]]]:
    number, issues = _decimal(value, sheet=sheet, row=row, column=column, field=field, max_value=Decimal("9999999999999.99"))
    if issues:
        return None, issues
    if number is None:
        return (None, []) if optional else (None, [_issue(sheet, row, column, "missing_required", f"{field}: el importe es obligatorio")])
    scaled = number * 100
    if scaled != scaled.to_integral_value():
        return None, [_issue(sheet, row, column, "fractional_cent", f"{field}: use como máximo dos decimales")]
    cents = int(scaled)
    if cents > 999_999_999_999_999:
        return None, [_issue(sheet, row, column, "money_range", f"{field}: excede el límite de 15 dígitos de centavos")]
    return cents, []


def _integer(value: Any, *, sheet: str, row: int, column: str, field: str) -> tuple[int | None, list[dict[str, Any]]]:
    if _blank(value):
        return None, [_issue(sheet, row, column, "missing_required", f"{field}: el valor es obligatorio")]
    if isinstance(value, bool):
        return None, [_issue(sheet, row, column, "invalid_integer", f"{field}: se esperaba un entero")]
    number, issues = _decimal(value, sheet=sheet, row=row, column=column, field=field, max_value=Decimal("2147483647"))
    if issues:
        return None, issues
    if number is None or number != number.to_integral_value():
        return None, [_issue(sheet, row, column, "invalid_integer", f"{field}: se esperaba un entero")]
    if number < 0:
        return None, [_issue(sheet, row, column, "invalid_integer", f"{field}: no puede ser negativo")]
    if number > 2_147_483_647:
        return None, [_issue(sheet, row, column, "integer_range", f"{field}: excede el límite entero permitido")]
    return int(number), []


def _enum(value: Any, allowed: Iterable[str], *, sheet: str, row: int, column: str, field: str, required: bool = True) -> tuple[str | None, list[dict[str, Any]]]:
    text = _text(value)
    if text is None:
        return (None, []) if not required else (None, [_issue(sheet, row, column, "missing_required", f"{field}: el valor es obligatorio")])
    normalized = text.casefold().replace(" ", "_")
    allowed = tuple(allowed)
    if normalized not in allowed:
        return None, [_issue(sheet, row, column, "invalid_enum", f"{field}: use uno de {', '.join(allowed)}")]
    return normalized, []


def _field_aliases(table: str) -> dict[str, str]:
    aliases: dict[str, str] = {}
    for field, values in TABLE_SPECS[table]["aliases"].items():
        for alias in (field, *values):
            aliases[_norm(alias)] = field
    return aliases


def _config_aliases() -> dict[str, str]:
    aliases: dict[str, str] = {}
    for field, values in CONFIG_ALIASES.items():
        for alias in (field, *values):
            aliases[_norm(alias)] = field
    return aliases


def _header_map(headers: Sequence[Any], table: str, sheet: str, row: int) -> tuple[dict[str, int], list[dict[str, Any]]]:
    aliases = _field_aliases(table)
    mapped: dict[str, int] = {}
    issues: list[dict[str, Any]] = []
    for index, header in enumerate(headers):
        key = aliases.get(_norm(header))
        if key is None:
            continue
        if key in mapped:
            issues.append(_issue(sheet, row, index + 1, "duplicate_column", f"La columna {key} aparece más de una vez"))
        else:
            mapped[key] = index
    for field in TABLE_SPECS[table]["required"]:
        if field not in mapped:
            issues.append(_issue(sheet, row, field, "missing_column", f"Falta la columna obligatoria {field}"))
    return mapped, issues


def _find_header(rows: Sequence[Sequence[Any]], table: str, sheet: str) -> tuple[int | None, dict[str, int], list[dict[str, Any]]]:
    candidates: list[tuple[int, dict[str, int], list[dict[str, Any]]]] = []
    for index, row in enumerate(rows[:30], start=1):
        if not any(not _blank(value) for value in row):
            continue
        mapped, issues = _header_map(row, table, sheet, index)
        if mapped:
            candidates.append((index, mapped, issues))
            if not issues:
                return index, mapped, []
    if candidates:
        return candidates[0]
    return None, {}, [_issue(sheet, 1, None, "missing_header", "No se encontró una fila de encabezados reconocible")]


def _row_values(row: Sequence[Any], mapping: Mapping[str, int]) -> dict[str, Any]:
    return {field: row[index] if index < len(row) else None for field, index in mapping.items()}


def _normalize_row(table: str, raw: Mapping[str, Any], *, sheet: str, row: int, as_of: datetime, assumptions: Counter[str]) -> tuple[dict[str, Any] | None, list[dict[str, Any]]]:
    issues: list[dict[str, Any]] = []
    spec = TABLE_SPECS[table]
    values = {field: raw.get(field) for field in spec["fields"]}
    for field in spec["required"]:
        if _blank(values.get(field)):
            issues.append(_issue(sheet, row, field, "missing_required", f"Falta el valor obligatorio {field}"))
    if issues:
        return None, issues

    result: dict[str, Any] = {}
    text_fields = {"order_id", "customer_ref", "invoice_id", "payment_id", "part_id", "description", "block_reason", "technician_ref"}
    for field in text_fields:
        if field not in values:
            continue
        value = _text(values[field])
        if value is not None:
            result[field] = value
        else:
            result[field] = None

    for field, allowed in spec["enums"].items():
        value, field_issues = _enum(values.get(field), allowed, sheet=sheet, row=row, column=field, field=field, required=field in spec["required"])
        result[field] = value
        issues.extend(field_issues)

    date_fields = {
        "orders": ("received_at", "delivered_at", "promised_at"),
        "invoices": ("issued_at", "due_at"),
        "payments": ("paid_at",),
        "inventory": (),
    }[table]
    for field in date_fields:
        value, field_issues, assumed = _parse_date(values.get(field), sheet=sheet, row=row, column=field, field=field)
        result[field] = value
        issues.extend(field_issues)
        if assumed and value is not None:
            assumptions["naive_excel_date_interpreted_as_utc"] += 1

    if table == "orders":
        value, field_issues = _decimal(values.get("estimated_hours"), sheet=sheet, row=row, column="estimated_hours", field="estimated_hours", max_value=Decimal("2147483647"))
        issues.extend(field_issues)
        if value is not None and value > Decimal("2147483647"):
            issues.append(_issue(sheet, row, "estimated_hours", "number_range", "estimated_hours excede el límite permitido"))
        result["estimated_hours"] = None if value is None else float(value)
    elif table == "invoices":
        result["amount_cents"], field_issues = _money_cents(values.get("amount_mxn"), sheet=sheet, row=row, column="amount_mxn", field="amount_mxn")
        issues.extend(field_issues)
    elif table == "payments":
        result["amount_cents"], field_issues = _money_cents(values.get("amount_mxn"), sheet=sheet, row=row, column="amount_mxn", field="amount_mxn")
        issues.extend(field_issues)
    elif table == "inventory":
        for field in ("on_hand", "reserved", "reorder_point"):
            result[field], field_issues = _integer(values.get(field), sheet=sheet, row=row, column=field, field=field)
            issues.extend(field_issues)
        result["unit_cost_cents"], field_issues = _money_cents(values.get("unit_cost_mxn"), sheet=sheet, row=row, column="unit_cost_mxn", field="unit_cost_mxn", optional=True)
        issues.extend(field_issues)

    # Order-local temporal checks are reliable even when there is no history.
    if table == "orders" and result.get("received_at") and result["received_at"] > _iso(as_of):
        issues.append(_issue(sheet, row, "received_at", "after_as_of", "received_at no puede estar después de Config.as_of"))
    if table == "orders" and result.get("delivered_at"):
        if result["delivered_at"] > _iso(as_of):
            issues.append(_issue(sheet, row, "delivered_at", "after_as_of", "delivered_at no puede estar después de Config.as_of"))
        if result.get("received_at") and result["delivered_at"] < result["received_at"]:
            issues.append(_issue(sheet, row, "delivered_at", "date_order", "delivered_at no puede ser anterior a received_at"))
    if table == "invoices" and result.get("issued_at") and result["issued_at"] > _iso(as_of):
        issues.append(_issue(sheet, row, "issued_at", "after_as_of", "issued_at no puede estar después de Config.as_of"))
    if table == "payments" and result.get("paid_at") and result["paid_at"] > _iso(as_of):
        issues.append(_issue(sheet, row, "paid_at", "after_as_of", "paid_at no puede estar después de Config.as_of"))

    if table == "orders" and result.get("status") == "delivered" and not result.get("delivered_at"):
        issues.append(_issue(sheet, row, "delivered_at", "required_for_status", "Una orden delivered requiere delivered_at"))
    if table == "invoices" and result.get("due_at") and result.get("issued_at") and result["due_at"] < result["issued_at"]:
        issues.append(_issue(sheet, row, "due_at", "date_order", "due_at no puede ser anterior a issued_at"))

    return result, issues


def _parse_bool(value: Any, *, sheet: str, row: int | None, column: str) -> tuple[bool | None, list[dict[str, Any]]]:
    if isinstance(value, bool):
        return value, []
    text = _text(value)
    if text is None:
        return None, [_issue(sheet, row, column, "missing_required", "synthetic es obligatorio")]
    normalized = text.casefold()
    if normalized in {"true", "1", "si", "sí", "yes"}:
        return True, []
    if normalized in {"false", "0", "no"}:
        return False, []
    return None, [_issue(sheet, row, column, "invalid_boolean", "synthetic debe ser true o false")]


def _config_from_rows(rows: Sequence[Sequence[Any]], sheet: str = CONFIG_SHEET) -> tuple[dict[str, Any], list[dict[str, Any]], int | None]:
    aliases = _config_aliases()
    issues: list[dict[str, Any]] = []
    values: dict[str, Any] = {}
    # Prefer a wide row if the workbook/CSV has columns named after config fields.
    for index, row in enumerate(rows[:20], start=1):
        mapping = {aliases.get(_norm(value)): col for col, value in enumerate(row) if aliases.get(_norm(value))}
        if {key for key in mapping if key} >= {"as_of", "snapshot_id", "synthetic"}:
            data_row = next((candidate for candidate in rows[index:] if any(not _blank(value) for value in candidate)), None)
            if data_row is None:
                data_row = []
            data = _row_values(data_row, mapping)
            values.update({key: data.get(key) for key in mapping})
            return values, issues, index

    header_index = None
    key_col = value_col = None
    for index, row in enumerate(rows[:20], start=1):
        normalized = [_norm(value) for value in row]
        if "campo" in normalized or "field" in normalized:
            header_index = index
            key_col = normalized.index("campo") if "campo" in normalized else normalized.index("field")
            if "valor" in normalized:
                value_col = normalized.index("valor")
            elif "value" in normalized:
                value_col = normalized.index("value")
            break
    if header_index is not None and value_col is not None and key_col is not None:
        for row_index, row in enumerate(rows[header_index:], start=header_index + 1):
            if key_col >= len(row) or _blank(row[key_col]):
                continue
            key = aliases.get(_norm(row[key_col]))
            if key is None:
                continue
            if key in values:
                issues.append(_issue(sheet, row_index, key_col + 1, "duplicate_config", f"Config.{key} aparece más de una vez"))
            else:
                values[key] = row[value_col] if value_col < len(row) else None
        return values, issues, header_index

    issues.append(_issue(sheet, 1, None, "missing_config", "Config requiere business_name, as_of, snapshot_id y synthetic"))
    return values, issues, None


def _source_records_from_csv(source: Path) -> tuple[dict[str, list[list[Any]]], list[dict[str, Any]], dict[str, str]]:
    records: dict[str, list[list[Any]]] = {}
    issues: list[dict[str, Any]] = []
    hashes: dict[str, str] = {}
    file_aliases = {_norm(name): table for table, name in SHEET_NAMES.items()}
    file_aliases.update({table: table for table in TABLES})
    file_aliases.update({"config": "config", "configuracion": "config"})
    files = [source] if source.is_file() else sorted(source.glob("*.csv"))
    if not files:
        return records, [_issue(str(source.name), None, None, "no_source_files", "No se encontraron archivos CSV")], hashes
    for path in files:
        hashes[path.name] = hashlib.sha256(path.read_bytes()).hexdigest()
        key = file_aliases.get(_norm(path.stem))
        if key is None:
            issues.append(_issue(path.name, 1, None, "unknown_table", "Nombre de archivo no reconocido"))
            continue
        if key in records:
            issues.append(_issue(path.name, 1, None, "duplicate_table", f"La tabla {key} aparece en más de un archivo reconocido"))
            continue
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                content = handle.read()
            dialect = csv.excel
            try:
                detected = csv.Sniffer().sniff(content[:4096], delimiters=",;\t|")
                # Sniffer can mistake a malformed row for a delimiter.  Keep
                # the result only for the delimiters the contract supports.
                if detected.delimiter in {",", ";", "\t", "|"}:
                    dialect = detected
            except csv.Error:
                pass
            rows = list(csv.reader(io.StringIO(content), dialect))
        except (OSError, UnicodeError, csv.Error) as exc:
            issues.append(_issue(path.name, None, None, "read_error", f"No se pudo leer el CSV: {exc}"))
            continue
        records[key] = rows
    return records, issues, hashes


def _source_records_from_workbook(source: Path) -> tuple[dict[str, list[list[Any]]], list[dict[str, Any]], dict[str, str]]:
    hashes = {source.name: hashlib.sha256(source.read_bytes()).hexdigest()}
    records: dict[str, list[list[Any]]] = {}
    issues: list[dict[str, Any]] = []
    try:
        book = openpyxl.load_workbook(source, data_only=False, read_only=True)
    except Exception as exc:  # openpyxl exposes several format-specific exceptions
        return records, [_issue(source.name, None, None, "read_error", f"No se pudo leer el libro: {exc}")], hashes
    name_aliases = {_norm(name): table for table, name in SHEET_NAMES.items()}
    name_aliases.update({table: table for table in TABLES})
    name_aliases.update({"config": "config", "configuracion": "config", _norm(CONFIG_SHEET): "config"})
    for worksheet in book.worksheets:
        key = name_aliases.get(_norm(worksheet.title))
        if key is None:
            if _norm(worksheet.title) not in IGNORED_SHEETS:
                issues.append(_issue(worksheet.title, 1, None, "unknown_sheet", "Hoja no reconocida; use solo las hojas del contrato"))
            continue
        if key in records:
            issues.append(_issue(worksheet.title, 1, None, "duplicate_table", f"La tabla {key} aparece en más de una hoja reconocida"))
            continue
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        records[key] = rows
    book.close()
    return records, issues, hashes


def _validate_references(tables: dict[str, list[dict[str, Any]]], sheets: Mapping[str, str]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    indexes = {
        "orders": {row["order_id"]: row for row in tables["orders"]},
        "invoices": {row["invoice_id"]: row for row in tables["invoices"]},
        "payments": {row["payment_id"]: row for row in tables["payments"]},
        "inventory": {row["part_id"]: row for row in tables["inventory"]},
    }
    for table, key in (("orders", "order_id"), ("invoices", "invoice_id"), ("payments", "payment_id"), ("inventory", "part_id")):
        seen: set[str] = set()
        for row in tables[table]:
            value = row.get(key)
            if value in seen:
                issues.append(_issue(sheets[table], row.get("__row"), key, "duplicate_ref", f"{key} duplicado: {value}"))
            seen.add(value)
    for row in tables["invoices"]:
        order_id = row.get("order_id")
        if order_id and order_id not in indexes["orders"]:
            issues.append(_issue(sheets["invoices"], row.get("__row"), "order_id", "bad_ref", f"No existe orders.order_id={order_id}"))
        if order_id and order_id in indexes["orders"] and row.get("customer_ref") != indexes["orders"][order_id].get("customer_ref"):
            issues.append(_issue(sheets["invoices"], row.get("__row"), "customer_ref", "ref_mismatch", "customer_ref no coincide con la orden vinculada"))
    for row in tables["payments"]:
        invoice_id = row.get("invoice_id")
        if invoice_id not in indexes["invoices"]:
            issues.append(_issue(sheets["payments"], row.get("__row"), "invoice_id", "bad_ref", f"No existe invoices.invoice_id={invoice_id}"))
        else:
            invoice = indexes["invoices"][invoice_id]
            if invoice.get("status") != "issued":
                issues.append(_issue(sheets["payments"], row.get("__row"), "invoice_id", "payment_on_void", f"No se aceptan pagos vinculados a factura {invoice_id} con estado void"))
            if row.get("paid_at") and invoice.get("issued_at") and row["paid_at"] < invoice["issued_at"]:
                issues.append(_issue(sheets["payments"], row.get("__row"), "paid_at", "date_order", "paid_at no puede ser anterior a issued_at"))
    paid_by_invoice: defaultdict[str, int] = defaultdict(int)
    for row in tables["payments"]:
        if row.get("invoice_id") in indexes["invoices"]:
            paid_by_invoice[row["invoice_id"]] += row.get("amount_cents") or 0
    for invoice_id, paid in paid_by_invoice.items():
        invoice = indexes["invoices"][invoice_id]
        if invoice.get("amount_cents") is not None and paid > invoice["amount_cents"]:
            issues.append(_issue(sheets["payments"], None, "amount_mxn", "payment_exceeds_invoice", f"Los pagos de {invoice_id} exceden el importe de la factura; créditos no están soportados"))
    return issues


def _quality_metadata(tables: dict[str, list[dict[str, Any]]], as_of: datetime) -> dict[str, Any]:
    order_ids = {row["order_id"] for row in tables["orders"]}
    invoiced_orders = {row["order_id"] for row in tables["invoices"] if row.get("order_id") and row.get("status") == "issued"}
    shortage = [row["part_id"] for row in tables["inventory"] if row["on_hand"] - row["reserved"] < 0]
    reorder = [row["part_id"] for row in tables["inventory"] if row["on_hand"] - row["reserved"] <= row["reorder_point"]]
    overdue_partial = []
    payments_by_invoice: defaultdict[str, int] = defaultdict(int)
    for payment in tables["payments"]:
        payments_by_invoice[payment["invoice_id"]] += payment["amount_cents"]
    for invoice in tables["invoices"]:
        if invoice.get("status") == "issued" and invoice.get("due_at") and invoice["due_at"] < _iso(as_of):
            paid = payments_by_invoice.get(invoice["invoice_id"], 0)
            if 0 < paid < invoice["amount_cents"]:
                overdue_partial.append(invoice["invoice_id"])
    return {
        "has_data": any(tables.values()),
        "counts": {name: len(rows) for name, rows in tables.items()},
        "cancelled_orders_excluded_from_open_work": sum(row.get("status") == "cancelled" for row in tables["orders"]),
        "delivered_orders_without_issued_invoice": sorted(row["order_id"] for row in tables["orders"] if row.get("status") == "delivered" and row["order_id"] not in invoiced_orders),
        "inventory_shortage_parts": sorted(shortage),
        "inventory_at_or_below_reorder_point": sorted(reorder),
        "overdue_partial_invoice_ids": sorted(overdue_partial),
        "settled_invoice_ids": sorted(
            invoice["invoice_id"] for invoice in tables["invoices"]
            if invoice.get("status") == "issued" and payments_by_invoice.get(invoice["invoice_id"], 0) == invoice["amount_cents"]
        ),
        "unknown_optional_values": {
            "delivered_at": sum(row.get("delivered_at") is None for row in tables["orders"]),
            "promised_at": sum(row.get("promised_at") is None for row in tables["orders"]),
            "block_reason": sum(row.get("block_reason") is None for row in tables["orders"]),
            "technician_ref": sum(row.get("technician_ref") is None for row in tables["orders"]),
            "estimated_hours": sum(row.get("estimated_hours") is None for row in tables["orders"]),
            "invoice_order_id": sum(row.get("order_id") is None for row in tables["invoices"]),
            "unit_cost_cents": sum(row.get("unit_cost_cents") is None for row in tables["inventory"]),
        },
        "cutoff": _iso(as_of),
    }


def load_client_input(source: str | Path, *, as_of: str | datetime | date | None = None) -> dict[str, Any]:
    """Load a client workbook or CSV directory into the canonical contract.

    ``as_of`` is an explicit override and must be timezone-aware when supplied.
    If omitted, ``Config.as_of`` is required.  Excel-native date cells are
    accepted as UTC by documented convention because Excel does not carry a
    timezone; ISO text with an offset remains the preferred format.
    """
    source = Path(source)
    if not source.exists():
        raise ClientInputError([_issue(str(source), None, None, "missing_source", "No existe la entrada indicada")])
    if source.is_dir() or source.suffix.casefold() == ".csv":
        records, source_issues, hashes = _source_records_from_csv(source)
    elif source.suffix.casefold() in {".xlsx", ".xlsm"}:
        records, source_issues, hashes = _source_records_from_workbook(source)
    else:
        raise ClientInputError([_issue(source.name, None, None, "unsupported_source", "Use un archivo .xlsx/.xlsm o una carpeta de CSV")])

    issues = list(source_issues)
    config_raw = records.get("config", [])
    config, config_issues, _ = _config_from_rows(config_raw)
    issues.extend(config_issues)
    explicit_as_of = None
    if as_of is not None:
        explicit_as_of, explicit_issues, _ = _as_explicit_as_of(as_of, sheet="API", row=None, column="as_of", allow_naive_excel=False)
        issues.extend(explicit_issues)
    config_as_of = None
    config_assumed = False
    if config.get("as_of") is not None:
        config_as_of, config_date_issues, config_assumed = _as_explicit_as_of(config.get("as_of"), sheet=CONFIG_SHEET, row=2, column="as_of")
        issues.extend(config_date_issues)
    cutoff = explicit_as_of or config_as_of
    if cutoff is None and not any(item["code"] == "missing_as_of" for item in issues):
        issues.append(_issue(CONFIG_SHEET, 2, "as_of", "missing_as_of", "Es necesario proporcionar un as_of explícito o Config.as_of"))
    if cutoff is None:
        cutoff = datetime.min.replace(tzinfo=UTC)

    snapshot_id = _text(config.get("snapshot_id"))
    business_name = _text(config.get("business_name"))
    if not business_name:
        issues.append(_issue(CONFIG_SHEET, 2, "business_name", "missing_required", "business_name es obligatorio"))
    if not snapshot_id:
        issues.append(_issue(CONFIG_SHEET, 2, "snapshot_id", "missing_required", "snapshot_id es obligatorio"))
    synthetic, synthetic_issues = _parse_bool(config.get("synthetic"), sheet=CONFIG_SHEET, row=2, column="synthetic")
    issues.extend(synthetic_issues)

    tables: dict[str, list[dict[str, Any]]] = {name: [] for name in TABLES}
    sheet_for_table = {name: SHEET_NAMES[name] for name in TABLES}
    assumptions: Counter[str] = Counter()
    if config_assumed:
        assumptions["naive_excel_date_interpreted_as_utc"] += 1
    missing_tables = [name for name in TABLES if name not in records]
    for table in missing_tables:
        issues.append(_issue(SHEET_NAMES[table], 1, None, "missing_sheet", f"Falta la hoja o archivo de {table}"))
    for table in TABLES:
        rows = records.get(table)
        if rows is None:
            continue
        header_row, mapping, header_issues = _find_header(rows, table, SHEET_NAMES[table])
        issues.extend(header_issues)
        if header_row is None:
            continue
        for row_number, row in enumerate(rows[header_row:], start=header_row + 1):
            if not any(not _blank(value) for value in row):
                continue
            # A formula cell is an input error even when data_only=False hides its cache.
            formula_columns = [index + 1 for index, value in enumerate(row) if isinstance(value, str) and value[:1] in {"=", "+", "@"}]
            if formula_columns:
                issues.append(_issue(SHEET_NAMES[table], row_number, formula_columns[0], "formula_input", "No se aceptan fórmulas en datos de entrada"))
                continue
            raw = _row_values(row, mapping)
            normalized, row_issues = _normalize_row(table, raw, sheet=SHEET_NAMES[table], row=row_number, as_of=cutoff, assumptions=assumptions)
            issues.extend(row_issues)
            if normalized is not None:
                normalized["__row"] = row_number
                tables[table].append(normalized)

    ref_issues = _validate_references(tables, sheet_for_table)
    issues.extend(ref_issues)
    if not any(tables.values()):
        issues.append(_issue(str(source.name), None, None, "no_data", "La entrada no contiene registros; no se puede emitir un reporte con ceros falsos"))
    if issues:
        raise ClientInputError(issues)

    for rows in tables.values():
        for row in rows:
            row.pop("__row", None)
    metadata = {
        "business_name": business_name,
        "as_of": _iso(cutoff),
        "snapshot_id": snapshot_id,
        "synthetic": synthetic,
        "source_kind": "workbook" if source.suffix.casefold() in {".xlsx", ".xlsm"} else "csv",
        "date_convention": "ISO-8601 con zona horaria; celdas Excel nativas sin zona se interpretan como UTC",
        "money_convention": "Decimal MXN exacto convertido a integer cents; máximo dos decimales",
        "quality": _quality_metadata(tables, cutoff),
        "assumptions": dict(assumptions),
    }
    return {"schema_version": SCHEMA_VERSION, "metadata": metadata, "tables": tables, "source_hashes": hashes}


def _write_data_sheet(book: xlsxwriter.Workbook, table: str, *, sample: bool, sample_rows: Sequence[Sequence[Any]]) -> None:
    worksheet = book.add_worksheet(SHEET_NAMES[table])
    worksheet.freeze_panes(4, 0)
    worksheet.set_zoom(90)
    title = book.add_format({"bold": True, "font_size": 16, "font_color": "#182B3A"})
    note = book.add_format({"text_wrap": True, "font_color": "#546975"})
    header = book.add_format({"bold": True, "font_color": "white", "bg_color": "#182B3A", "text_wrap": True})
    worksheet.merge_range(0, 0, 0, len(TABLE_SPECS[table]["headers"]) - 1, f"Milenio · {SHEET_NAMES[table]}", title)
    worksheet.merge_range(1, 0, 1, len(TABLE_SPECS[table]["headers"]) - 1, "Capture una fila por registro. Fechas preferentemente ISO-8601 con zona horaria; los blancos son desconocidos y no significan cero.", note)
    start = 3
    worksheet.write_row(start, 0, TABLE_SPECS[table]["headers"], header)
    rows = list(sample_rows) if sample else []
    if rows:
        worksheet.add_table(start, 0, start + len(rows), len(TABLE_SPECS[table]["headers"]) - 1, {
            "name": f"Tabla_{table}",
            "style": "Table Style Medium 2",
            "columns": [{"header": value} for value in TABLE_SPECS[table]["headers"]],
            "data": rows,
        })
    else:
        # Keep a data-entry row visible without inventing a record.
        worksheet.write_row(start + 1, 0, [None] * len(TABLE_SPECS[table]["headers"]))
    for index, field in enumerate(TABLE_SPECS[table]["fields"]):
        worksheet.set_column(index, index, 20 if field.endswith("_at") or field.endswith("_id") else 18)
    for field, allowed in TABLE_SPECS[table]["enums"].items():
        col = TABLE_SPECS[table]["fields"].index(field)
        worksheet.data_validation(start + 1, col, start + 500, col, {"validate": "list", "source": list(allowed), "input_title": field, "input_message": "Elija un valor de la lista"})
    for field in ("received_at", "delivered_at", "promised_at", "issued_at", "due_at", "paid_at"):
        if field in TABLE_SPECS[table]["fields"]:
            col = TABLE_SPECS[table]["fields"].index(field)
            worksheet.data_validation(start + 1, col, start + 500, col, {"validate": "custom", "value": f'=OR({xlsxwriter.utility.xl_col_to_name(col)}5="",ISNUMBER({xlsxwriter.utility.xl_col_to_name(col)}5),ISTEXT({xlsxwriter.utility.xl_col_to_name(col)}5))', "error_message": "Use ISO-8601 con zona horaria o fecha Excel", "error_type": "stop"})


def _sample_rows(as_of: datetime) -> dict[str, list[list[Any]]]:
    day = as_of.astimezone(UTC).date()
    def iso(days: int, hour: int = 12) -> str:
        return _iso(datetime.combine(day, time(hour), tzinfo=UTC) + timedelta(days=days))

    # Rows intentionally cover the client-facing cases documented in the contract.
    return {
        "orders": [
            ["ORD-001", "CUST-001", "particular", iso(-12), "", "waiting_parts", iso(-8), "Filtro pendiente", "TECH-01", 3.5],
            ["ORD-002", "CUST-002", "fleet", iso(-10), "", "in_service", iso(-4), "", "TECH-02", 5],
            ["ORD-003", "CUST-001", "particular", iso(-9), "", "ready", iso(-2), "", "TECH-01", 2],
            ["ORD-004", "CUST-003", "particular", iso(-15), iso(-7), "delivered", iso(-10), "", "TECH-03", 4],
            ["ORD-005", "CUST-004", "fleet", iso(-20), "", "cancelled", iso(-18), "Cliente canceló", "", ""],
            ["ORD-006", "CUST-005", "fleet", iso(-18), iso(-12), "delivered", iso(-14), "", "TECH-02", 6],
        ],
        "invoices": [
            ["INV-001", "ORD-002", "CUST-002", iso(-9), iso(-2), "1500.00", "issued"],
            ["INV-002", "ORD-006", "CUST-005", iso(-11), iso(-1), "800.00", "issued"],
        ],
        "payments": [
            ["PAY-001", "INV-001", iso(-1), "500.00", "transfer"],
            ["PAY-002", "INV-002", iso(-5), "800.00", "card"],
        ],
        "inventory": [
            ["PART-001", "Filtro de aceite", 1, 3, 2, "250.00"],
            ["PART-002", "Bujía", 20, 2, 5, "90.00"],
        ],
    }


def write_client_template(destination: str | Path, *, sample: bool = False, as_of: str | datetime | None = None, business_name: str = "", snapshot_id: str | None = None, synthetic: bool = True) -> Path:
    """Write a blank or synthetic workbook template.

    A sample requires a caller-provided cutoff; this avoids silently reusing
    the project's synthetic demo clock in a client-facing file.
    """
    destination = Path(destination)
    if type(synthetic) is not bool:
        raise ValueError("synthetic debe ser true o false")
    if sample and synthetic is not True:
        raise ValueError("los datos de muestra siempre deben ser synthetic=true")
    if as_of is None:
        if sample:
            raise ValueError("sample templates require an explicit timezone-aware as_of")
        cutoff = None
    else:
        cutoff, issues, _ = _as_explicit_as_of(as_of, sheet="API", column="as_of", allow_naive_excel=False)
        if issues or cutoff is None:
            raise ValueError(issues[0]["message"] if issues else "as_of inválido")
    if sample and not _text(business_name):
        business_name = "Ejemplo sintético"
    destination.parent.mkdir(parents=True, exist_ok=True)
    book = xlsxwriter.Workbook(destination, {"strings_to_formulas": False, "strings_to_urls": False, "strings_to_numbers": False})
    book.set_properties({"title": "Milenio · Plantilla de entrada local", "comments": "Datos sintéticos de ejemplo; no es un conector de negocio."})
    title = book.add_format({"bold": True, "font_size": 18, "font_color": "#182B3A"})
    note = book.add_format({"text_wrap": True, "font_color": "#546975"})
    config = book.add_worksheet(CONFIG_SHEET)
    config.hide_gridlines(2)
    config.merge_range("A1:C1", "Milenio · Configuración del corte", title)
    config.write_row("A3", ["campo", "valor", "nota"], book.add_format({"bold": True, "font_color": "white", "bg_color": "#182B3A"}))
    config_rows = [
        ["business_name", business_name, "Obligatorio; no incluya nombres de personas ni datos de contacto."],
        ["as_of", _iso(cutoff) if cutoff else "", "Obligatorio. ISO-8601 con zona horaria; ejemplo configurable, no reloj DEMO."],
        ["snapshot_id", snapshot_id or ("sample-client-001" if sample else ""), "Identificador estable del corte."],
        ["synthetic", "true" if synthetic else "false", "Use true solo para datos fabricados."],
    ]
    config.add_table(2, 0, 2 + len(config_rows), 2, {"name": "Tabla_Config", "style": "Table Style Medium 2", "columns": [{"header": "campo"}, {"header": "valor"}, {"header": "nota"}], "data": config_rows})
    config.set_column("A:A", 20); config.set_column("B:B", 34); config.set_column("C:C", 72)
    config.data_validation("B7", {"validate": "list", "source": ["true", "false"], "input_title": "synthetic", "input_message": "Elija true o false"})

    instructions = book.add_worksheet(INSTRUCTIONS_SHEET)
    instructions.hide_gridlines(2)
    instructions.set_column("A:A", 28); instructions.set_column("B:B", 110)
    instructions.merge_range("A1:B1", "Cómo preparar la entrada", title)
    instruction_rows = [
        ("Propósito", "Este libro sirve para análisis local por corte. No envía mensajes, no agenda, no cobra y no modifica sistemas del negocio."),
        ("Fechas", "Prefiera texto ISO-8601 con zona horaria, por ejemplo 2026-09-22T12:00:00-07:00. Una celda de fecha Excel sin zona se interpreta como UTC y queda registrada como supuesto."),
        ("Importes", "Use números MXN con máximo dos decimales. El adaptador convierte Decimal exacto a integer cents; no use NaN, infinito ni fórmulas."),
        ("Blancos", "Un blanco significa desconocido o no informado y se conserva como null. No se convierte en cero."),
        ("Referencias", "invoice.order_id, invoice.customer_ref y payment.invoice_id se reconcilian con las tablas disponibles. Se rechazan duplicados y referencias inexistentes."),
        ("Calidad", "La carga se detiene si falta estructura, hay fechas después de as_of o importes inválidos. Escasez, vencidos parciales, cancelados y entregados sin factura quedan como señales explícitas."),
    ]
    instructions.add_table(2, 0, 2 + len(instruction_rows), 1, {"name": "Tabla_Instrucciones", "style": "Table Style Medium 2", "columns": [{"header": "Tema"}, {"header": "Regla"}], "data": [list(row) for row in instruction_rows]})

    sample_rows = _sample_rows(cutoff) if sample and cutoff else {table: [] for table in TABLES}
    for table in TABLES:
        _write_data_sheet(book, table, sample=sample, sample_rows=sample_rows[table])
    book.close()
    return destination


def write_client_csv_templates(directory: str | Path, *, sample: bool = False, as_of: str | datetime | None = None, business_name: str = "", snapshot_id: str | None = None, synthetic: bool = True) -> Path:
    """Write a directory of CSV equivalents for batch clients."""
    directory = Path(directory)
    if type(synthetic) is not bool:
        raise ValueError("synthetic debe ser true o false")
    if sample and synthetic is not True:
        raise ValueError("los datos de muestra siempre deben ser synthetic=true")
    directory.mkdir(parents=True, exist_ok=True)
    if sample and as_of is None:
        raise ValueError("sample templates require an explicit timezone-aware as_of")
    cutoff = None
    if as_of is not None:
        cutoff, issues, _ = _as_explicit_as_of(as_of, sheet="API", column="as_of", allow_naive_excel=False)
        if issues or cutoff is None:
            raise ValueError(issues[0]["message"] if issues else "as_of inválido")
    if sample and not _text(business_name):
        business_name = "Ejemplo sintético"
    rows = _sample_rows(cutoff) if sample and cutoff else {table: [] for table in TABLES}
    with (directory / "config.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["campo", "valor"])
        writer.writerows([
            ["business_name", business_name],
            ["as_of", _iso(cutoff) if cutoff else ""],
            ["snapshot_id", snapshot_id or ("sample-client-001" if sample else "")],
            ["synthetic", "true" if synthetic else "false"],
        ])
    for table in TABLES:
        with (directory / f"{table}.csv").open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(TABLE_SPECS[table]["headers"])
            writer.writerows(rows[table])
    return directory


def generate_client_templates(directory: str | Path, *, as_of: str | datetime, business_name: str = "", snapshot_id: str = "sample-client-001") -> dict[str, Path]:
    """Generate blank and sample workbook/CSV bundles from one explicit cutoff."""
    directory = Path(directory)
    blank_xlsx = write_client_template(directory / "client_input_blank.xlsx", sample=False, business_name=business_name, snapshot_id="", synthetic=True)
    sample_xlsx = write_client_template(directory / "client_input_sample.xlsx", sample=True, as_of=as_of, business_name=business_name, snapshot_id=snapshot_id, synthetic=True)
    blank_csv = write_client_csv_templates(directory / "blank_csv", sample=False, business_name=business_name, snapshot_id="", synthetic=True)
    sample_csv = write_client_csv_templates(directory / "sample_csv", sample=True, as_of=as_of, business_name=business_name, snapshot_id=snapshot_id, synthetic=True)
    return {"blank_workbook": blank_xlsx, "sample_workbook": sample_xlsx, "blank_csv": blank_csv, "sample_csv": sample_csv}


__all__ = [
    "ClientInputError",
    "SCHEMA_VERSION",
    "TABLES",
    "TABLE_SPECS",
    "generate_client_templates",
    "load_client_input",
    "write_client_csv_templates",
    "write_client_template",
]
