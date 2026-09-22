"""Offline follow-through for evidence-backed human action reviews.

This module deliberately stops at a local Excel handoff.  It does not contact a
customer, update a CRM, send a message, schedule work, or perform any other
business operation.  Decision rows are written from a synthetic snapshot and
the editable columns are annotations reported by the person using the sheet.
"""

from __future__ import annotations

import json
import os
import tempfile
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping

import openpyxl
import xlsxwriter


class ClientActionError(ValueError):
    """Raised when a decision workbook or review import is invalid."""


REVIEW_STATUSES = ("pending", "accepted", "in_progress", "done", "dismissed")

# These names are intentionally stable machine-readable fields.  The workbook
# uses the Spanish labels below while the importer maps them back to these keys.
WORKBOOK_COLUMNS = (
    "action_id",
    "category",
    "title",
    "priority",
    "owner_role",
    "why_now",
    "recommended_next_step",
    "amount_at_risk_cents",
    "source_ids",
    "evidence",
    "source_snapshot",
    "owner",
    "status",
    "target_date",
    "note",
    "outcome_evidence",
)

WORKBOOK_HEADERS = {
    "action_id": "ID de acción",
    "category": "Categoría",
    "title": "Título",
    "priority": "Prioridad",
    "owner_role": "Rol sugerido",
    "why_now": "Por qué ahora",
    "recommended_next_step": "Siguiente paso recomendado",
    "amount_at_risk_cents": "Importe relacionado (MXN)",
    "source_ids": "IDs de fuente",
    "evidence": "Evidencia de origen",
    "source_snapshot": "Snapshot de origen",
    "owner": "Responsable",
    "status": "Estado",
    "target_date": "Fecha objetivo",
    "note": "Nota de revisión",
    "outcome_evidence": "Evidencia del resultado",
}
WORKBOOK_HEADER_ROW = tuple(WORKBOOK_HEADERS[field] for field in WORKBOOK_COLUMNS)
_HEADER_FIELDS = {header: field for field, header in WORKBOOK_HEADERS.items()}

ORIGIN_COLUMNS = WORKBOOK_COLUMNS[:11]
REVIEW_COLUMNS = WORKBOOK_COLUMNS[11:]
REQUIRED_DECISION_FIELDS = (
    "action_id",
    "category",
    "title",
    "priority",
    "owner_role",
    "why_now",
    "recommended_next_step",
    "amount_at_risk_cents",
    "evidence",
    "source_ids",
)


def _error(message: str) -> None:
    raise ClientActionError(message)


def _text(value: Any, field: str, *, required: bool = True) -> str | None:
    if value is None:
        if required:
            _error(f"{field} is required")
        return None
    if not isinstance(value, str):
        _error(f"{field} must be text")
    value = value.strip()
    if required and not value:
        _error(f"{field} is required")
    return value or None


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _json(value: Any) -> str:
    try:
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    except (TypeError, ValueError) as exc:
        raise ClientActionError("decision contains a value that cannot be serialized") from exc


def _parse_json(value: Any, field: str) -> Any:
    if not isinstance(value, str) or not value.strip():
        _error(f"{field} must contain JSON")
    try:
        return json.loads(value)
    except (TypeError, ValueError) as exc:
        raise ClientActionError(f"{field} contains invalid JSON") from exc


def _validate_metadata(metadata: Mapping[str, Any]) -> dict[str, Any]:
    if not isinstance(metadata, Mapping):
        _error("metadata must be an object")
    snapshot_id = _text(metadata.get("snapshot_id"), "metadata.snapshot_id")
    as_of = _text(metadata.get("as_of"), "metadata.as_of")
    if type(metadata.get("synthetic")) is not bool:
        _error("metadata.synthetic must be boolean")
    return {**dict(metadata), "snapshot_id": snapshot_id, "as_of": as_of, "synthetic": metadata["synthetic"]}


def _validate_decisions(decisions: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    if isinstance(decisions, (str, bytes)):
        _error("decisions must be a sequence of objects")
    try:
        values = list(decisions)
    except TypeError as exc:
        raise ClientActionError("decisions must be a sequence of objects") from exc

    normalized: list[dict[str, Any]] = []
    ids: set[str] = set()
    for index, decision in enumerate(values, start=1):
        if not isinstance(decision, Mapping):
            _error(f"decision {index} must be an object")
        missing = [field for field in REQUIRED_DECISION_FIELDS if field not in decision]
        if missing:
            _error(f"decision {index} is missing: {', '.join(missing)}")
        item = dict(decision)
        for field in ("action_id", "category", "title", "priority", "owner_role", "why_now", "recommended_next_step"):
            item[field] = _text(item.get(field), f"decision {index}.{field}")
        action_id = item["action_id"]
        if action_id in ids:
            _error(f"duplicate action_id: {action_id}")
        ids.add(action_id)

        amount = item["amount_at_risk_cents"]
        if amount is not None and (type(amount) is not int or amount < 0):
            _error(f"decision {index}.amount_at_risk_cents must be a non-negative integer or null")
        if not isinstance(item["evidence"], list):
            _error(f"decision {index}.evidence must be a list")
        if not all(isinstance(ref, Mapping) for ref in item["evidence"]):
            _error(f"decision {index}.evidence must contain objects")
        if not isinstance(item["source_ids"], list):
            _error(f"decision {index}.source_ids must be a list")
        if not all(isinstance(source_id, str) and source_id.strip() for source_id in item["source_ids"]):
            _error(f"decision {index}.source_ids must contain non-empty text")

        # Validate serializability now so writing cannot create a half-valid
        # workbook after an earlier row has already been prepared.
        _json(item["source_ids"])
        _json(item["evidence"])
        normalized.append(item)
    return normalized


def _normalise_amount(value: Any, field: str) -> int | None:
    if _blank(value):
        return None
    if isinstance(value, bool):
        _error(f"{field} must be a non-negative MXN amount or blank")
    try:
        amount_mxn = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        _error(f"{field} must be a non-negative MXN amount or blank")
    cents = amount_mxn * 100
    if not cents.is_finite() or cents < 0 or cents != cents.to_integral_value():
        _error(f"{field} must represent whole MXN cents")
    return int(cents)


def _normalise_date(value: Any, field: str) -> str | None:
    if _blank(value):
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        if value:
            return value
    _error(f"{field} must be a date or blank")


def _compare_origin(row: Mapping[str, Any], expected: Mapping[str, Any], snapshot_id: str) -> None:
    if row["source_snapshot"] != snapshot_id:
        _error(f"source_snapshot mismatch for action_id {row['action_id']}")
    comparisons = {
        "action_id": expected["action_id"],
        "category": expected["category"],
        "title": expected["title"],
        "priority": expected["priority"],
        "owner_role": expected["owner_role"],
        "why_now": expected["why_now"],
        "recommended_next_step": expected["recommended_next_step"],
        "amount_at_risk_cents": expected["amount_at_risk_cents"],
        "source_ids": expected["source_ids"],
        "evidence": expected["evidence"],
    }
    for field, expected_value in comparisons.items():
        if row[field] != expected_value:
            _error(f"immutable field changed for action_id {row['action_id']}: {field}")


def write_action_workbook(
    path: str | Path,
    decisions: Iterable[Mapping[str, Any]],
    metadata: Mapping[str, Any],
) -> dict[str, Any]:
    """Write a two-sheet Spanish action handoff for one synthetic snapshot.

    All rows start as ``pending`` and all human fields are blank.  Writing the
    workbook does not imply acceptance, approval, or execution.
    """

    target = Path(path)
    if target.exists() and target.is_dir():
        _error("workbook path must be a file")
    target.parent.mkdir(parents=True, exist_ok=True)
    meta = _validate_metadata(metadata)
    rows = _validate_decisions(decisions)

    book = xlsxwriter.Workbook(str(target), {"strings_to_formulas": False, "strings_to_urls": False})
    try:
        book.set_properties(
            {
                "title": "Milenio | Seguimiento de acciones",
                "subject": "Revisión humana offline de decisiones del corte",
                "author": "Taller Milenio Analytics",
                "comments": "Entrega local; no ejecuta acciones de negocio.",
            }
        )
        navy, ink = "#182B3A", "#263238"
        title = book.add_format({"font_name": "Aptos", "font_size": 18, "bold": True, "font_color": navy})
        subtitle = book.add_format({"font_name": "Aptos", "font_size": 10, "font_color": "#546975", "text_wrap": True})
        header = book.add_format({"font_name": "Aptos", "bold": True, "font_color": "white", "bg_color": navy, "text_wrap": True, "valign": "vcenter"})
        text = book.add_format({"font_name": "Aptos", "font_size": 10, "font_color": ink, "valign": "top", "text_wrap": True})
        money = book.add_format({"font_name": "Aptos", "font_size": 10, "font_color": ink, "num_format": '$#,##0.00;($#,##0.00);"—"', "valign": "top"})
        status = book.add_format({"font_name": "Aptos", "font_size": 10, "font_color": ink, "bg_color": "#FFF2CC", "valign": "top"})
        blank = book.add_format({"font_name": "Aptos", "font_size": 10, "font_color": "#7A868C", "bg_color": "#FAFAFA", "valign": "top", "text_wrap": True})

        sheet = book.add_worksheet("Seguimiento")
        sheet.hide_gridlines(2)
        sheet.freeze_panes(1, 3)
        sheet.set_zoom(85)
        sheet.set_row(0, 32)
        widths = {
            "action_id": 17,
            "category": 16,
            "title": 30,
            "priority": 12,
            "owner_role": 19,
            "why_now": 34,
            "recommended_next_step": 38,
            "amount_at_risk_cents": 18,
            "source_ids": 26,
            "evidence": 44,
            "source_snapshot": 22,
            "owner": 22,
            "status": 16,
            "target_date": 17,
            "note": 34,
            "outcome_evidence": 38,
        }
        hidden_origin = set(ORIGIN_COLUMNS) - {"title", "priority"}
        for column, name in enumerate(WORKBOOK_COLUMNS):
            options = {"hidden": True} if name in hidden_origin else None
            sheet.set_column(column, column, widths[name], money if name == "amount_at_risk_cents" else text, options)
        sheet.write_row(0, 0, WORKBOOK_HEADER_ROW, header)

        for row_number, decision in enumerate(rows, start=1):
            sheet.set_row(row_number, 56)
            values = [
                decision["action_id"],
                decision["category"],
                decision["title"],
                decision["priority"],
                decision["owner_role"],
                decision["why_now"],
                decision["recommended_next_step"],
                None if decision["amount_at_risk_cents"] is None else decision["amount_at_risk_cents"] / 100,
                _json(decision["source_ids"]),
                _json(decision["evidence"]),
                meta["snapshot_id"],
                "",
                "pending",
                "",
                "",
                "",
            ]
            for column, value in enumerate(values):
                cell_format = status if WORKBOOK_COLUMNS[column] == "status" else blank if WORKBOOK_COLUMNS[column] in REVIEW_COLUMNS and not value else money if WORKBOOK_COLUMNS[column] == "amount_at_risk_cents" else text
                if value is None:
                    sheet.write_blank(row_number, column, None, cell_format)
                elif isinstance(value, (int, float)) and not isinstance(value, bool):
                    sheet.write_number(row_number, column, value, cell_format)
                else:
                    sheet.write(row_number, column, value, cell_format)
        if rows:
            sheet.add_table(0, 0, len(rows), len(WORKBOOK_COLUMNS) - 1, {
                "name": "ActionFollowThrough",
                "style": "Table Style Medium 2",
                "columns": [{"header": field} for field in WORKBOOK_HEADER_ROW],
            })
        sheet.data_validation(1, WORKBOOK_COLUMNS.index("status"), max(1, len(rows)), WORKBOOK_COLUMNS.index("status"), {
            "validate": "list",
            "source": list(REVIEW_STATUSES),
            "input_title": "Estado de revisión",
            "input_message": "Seleccione un estado; no ejecuta ninguna acción externa.",
            "error_title": "Estado inválido",
            "error_message": "Use pending, accepted, in_progress, done o dismissed.",
        })
        if rows:
            status_column = WORKBOOK_COLUMNS.index("status")
            sheet.conditional_format(1, status_column, len(rows), status_column, {
                "type": "text", "criteria": "containing", "value": "done",
                "format": book.add_format({"bg_color": "#E2F0D9", "font_color": "#375623"}),
            })

        instructions = book.add_worksheet("Instrucciones")
        instructions.hide_gridlines(2)
        instructions.set_column("A:A", 24)
        instructions.set_column("B:B", 110)
        instructions.merge_range("A1:B2", "MILENIO / Seguimiento de acciones", title)
        data_label = "sintéticos" if meta["synthetic"] else "del cliente y confidenciales"
        instructions.merge_range("A3:B4", f"Libro local para que una persona registre el seguimiento de decisiones con evidencia. Los datos son {data_label} y se refieren exclusivamente al corte indicado.", subtitle)
        instruction_rows = [
            ("snapshot_id", meta["snapshot_id"]),
            ("as_of", meta["as_of"]),
            ("synthetic", "true" if meta["synthetic"] else "false"),
            ("Qué se puede editar", "Responsable, Estado, Fecha objetivo, Nota de revisión y Evidencia del resultado en la hoja Seguimiento."),
            ("Qué no se puede editar", "La sección de origen se conserva para auditoría. Los campos técnicos de IDs, evidencia, snapshot e importe relacionado no deben cambiarse; si no coinciden con el origen, la importación se rechaza."),
            ("Contexto de origen", "El motivo, el siguiente paso y la evidencia completa también aparecen en INICIO.html y Gerencia.xlsx. Las columnas de origen de Seguimiento están ocultas para facilitar la reunión; puede mostrarlas cuando necesite auditar el vínculo."),
            ("Para marcar done", "Escriba Responsable, Nota de revisión y Evidencia del resultado. El resultado es auto-reportado; la hoja no prueba por sí sola el resultado del negocio."),
            ("Estados permitidos", "Use los códigos: pending = pendiente; accepted = aceptada; in_progress = en curso; done = cerrada con evidencia; dismissed = descartada. accepted e in_progress requieren responsable y fecha objetivo; dismissed requiere nota."),
            ("Límite operativo", "Este libro no contacta clientes, no cambia un CRM, no agenda, no cobra, no compra y no publica. Toda acción de negocio requiere un proceso humano separado."),
            ("Auditoría", "Conserve el snapshot original. El registro JSONL de importación debe incluir un identificador de fuente y un revisor declarado por quien lo registra; esa etiqueta es auto-reportada y no verifica una identidad."),
        ]
        for row_number, (label, value) in enumerate(instruction_rows, start=5):
            instructions.write(row_number - 1, 0, label, header if row_number == 5 else text)
            instructions.write(row_number - 1, 1, value, text)
        instructions.set_row(8, 46)
        instructions.set_row(9, 58)
        instructions.set_row(10, 48)
        instructions.set_row(12, 58)
    finally:
        book.close()

    return {
        "path": str(target),
        "rows": len(rows),
        "worksheets": ["Seguimiento", "Instrucciones"],
        "snapshot_id": meta["snapshot_id"],
        "as_of": meta["as_of"],
        "synthetic": meta["synthetic"],
        "external_business_action": False,
    }


def load_action_reviews(
    workbook: str | Path,
    expected_decisions: Iterable[Mapping[str, Any]],
    metadata: Mapping[str, Any],
) -> list[dict[str, Any]]:
    """Load and validate self-reported review fields from a local workbook."""

    source = Path(workbook)
    if not source.exists() or not source.is_file():
        _error("workbook does not exist")
    meta = _validate_metadata(metadata)
    expected = _validate_decisions(expected_decisions)
    by_id = {decision["action_id"]: decision for decision in expected}

    try:
        book = openpyxl.load_workbook(source, read_only=True, data_only=False)
    except Exception as exc:  # openpyxl exposes several format-specific errors.
        raise ClientActionError("could not open action workbook") from exc
    try:
        if "Seguimiento" not in book.sheetnames:
            _error("workbook is missing the Seguimiento sheet")
        sheet = book["Seguimiento"]
        try:
            headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        except StopIteration:
            _error("Seguimiento sheet is empty")
        if len(headers) != len(WORKBOOK_COLUMNS) or set(headers) != set(WORKBOOK_HEADER_ROW):
            _error("Seguimiento headers do not match the action workbook contract")
        header_fields = [_HEADER_FIELDS.get(header) for header in headers]
        if any(field is None for field in header_fields) or len(set(header_fields)) != len(WORKBOOK_COLUMNS):
            _error("Seguimiento headers contain unknown or duplicate labels")

        seen: set[str] = set()
        reviews: list[dict[str, Any]] = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2, values_only=False), start=2):
            cells = list(cells[: len(WORKBOOK_COLUMNS)])
            values_by_field = {field: cell.value for field, cell in zip(header_fields, cells)}
            values = [values_by_field[field] for field in WORKBOOK_COLUMNS]
            if all(_blank(value) for value in values):
                continue
            if len(values) != len(WORKBOOK_COLUMNS):
                _error(f"row {row_number} has the wrong number of columns")
            for field, cell in zip(header_fields, cells):
                if field in REVIEW_COLUMNS and cell.data_type == "f":
                    _error(f"formula cells are not allowed in review field {field}: row {row_number}")
            action_id = _text(values_by_field["action_id"], f"row {row_number}.action_id")
            if action_id in seen:
                _error(f"duplicate action_id in workbook: {action_id}")
            seen.add(action_id)
            if action_id not in by_id:
                _error(f"unknown action_id in workbook: {action_id}")
            expected_decision = by_id[action_id]
            source_ids = _parse_json(values_by_field["source_ids"], f"row {row_number}.source_ids")
            evidence = _parse_json(values_by_field["evidence"], f"row {row_number}.evidence")
            row = {
                "action_id": action_id,
                "category": _text(values_by_field["category"], f"row {row_number}.category"),
                "title": _text(values_by_field["title"], f"row {row_number}.title"),
                "priority": _text(values_by_field["priority"], f"row {row_number}.priority"),
                "owner_role": _text(values_by_field["owner_role"], f"row {row_number}.owner_role"),
                "why_now": _text(values_by_field["why_now"], f"row {row_number}.why_now"),
                "recommended_next_step": _text(values_by_field["recommended_next_step"], f"row {row_number}.recommended_next_step"),
                "amount_at_risk_cents": _normalise_amount(values_by_field["amount_at_risk_cents"], f"row {row_number}.amount_at_risk_cents"),
                "source_ids": source_ids,
                "evidence": evidence,
                "source_snapshot": _text(values_by_field["source_snapshot"], f"row {row_number}.source_snapshot"),
            }
            _compare_origin(row, expected_decision, meta["snapshot_id"])

            owner = _text(values_by_field["owner"], f"row {row_number}.owner", required=False)
            status = _text(values_by_field["status"], f"row {row_number}.status")
            if status not in REVIEW_STATUSES:
                _error(f"illegal status for action_id {action_id}: {status}")
            target_date = _normalise_date(values_by_field["target_date"], f"row {row_number}.target_date")
            note = _text(values_by_field["note"], f"row {row_number}.note", required=False)
            outcome_evidence = _text(values_by_field["outcome_evidence"], f"row {row_number}.outcome_evidence", required=False)
            if status in {"accepted", "in_progress"} and not all((owner, target_date)):
                _error(f"{status} review requires owner and target_date: {action_id}")
            if status == "dismissed" and not note:
                _error(f"dismissed review requires note: {action_id}")
            if status == "done" and not all((owner, note, outcome_evidence)):
                _error(f"done review requires owner, note and outcome_evidence: {action_id}")
            reviews.append({
                "action_id": action_id,
                "source_snapshot": row["source_snapshot"],
                "owner": owner,
                "status": status,
                "target_date": target_date,
                "note": note,
                "outcome_evidence": outcome_evidence,
                "self_reported": True,
                "external_business_action": False,
            })
        missing = set(by_id) - seen
        if missing:
            _error("workbook is missing action_id(s): " + ", ".join(sorted(missing)))
        return reviews
    finally:
        book.close()


def append_review_log(
    output: str | Path,
    reviews: Iterable[Mapping[str, Any]],
    source_hash: str | None = None,
    reviewer: str | None = None,
    *,
    sourcehash: str | None = None,
) -> dict[str, Any]:
    """Atomically append one review import event to a JSONL audit file.

    ``reviewer`` is a required self-reported label.  It is stored as such and
    never interpreted as proof of a person's identity or authority.  A source
    hash already present in the log is an idempotent no-op.
    """

    if source_hash is not None and sourcehash is not None:
        _error("provide source_hash or sourcehash, not both")
    source_hash = source_hash if source_hash is not None else sourcehash
    source_hash = _text(source_hash, "source_hash")
    reviewer = _text(reviewer, "reviewer")
    if isinstance(reviews, (str, bytes)):
        _error("reviews must be a sequence of objects")
    try:
        review_values = list(reviews)
    except TypeError as exc:
        raise ClientActionError("reviews must be a sequence of objects") from exc
    serializable_reviews: list[dict[str, Any]] = []
    for index, review in enumerate(review_values, start=1):
        if not isinstance(review, Mapping):
            _error(f"review {index} must be an object")
        item = dict(review)
        _text(item.get("action_id"), f"review {index}.action_id")
        status = _text(item.get("status"), f"review {index}.status")
        if status not in REVIEW_STATUSES:
            _error(f"illegal status for review {index}: {status}")
        # Keep the imported record intact but force the boundary marker so an
        # audit line cannot be mistaken for an external business write.
        item["self_reported"] = True
        item["external_business_action"] = False
        _json(item)
        serializable_reviews.append(item)

    target = Path(output)
    if target.exists() and target.is_dir():
        _error("audit output must be a file")
    target.parent.mkdir(parents=True, exist_ok=True)
    existing = ""
    if target.exists():
        try:
            existing = target.read_text(encoding="utf-8")
        except OSError as exc:
            raise ClientActionError("could not read existing audit log") from exc
        for line_number, line in enumerate(existing.splitlines(), start=1):
            if not line.strip():
                continue
            try:
                prior = json.loads(line)
            except (TypeError, ValueError) as exc:
                raise ClientActionError(f"audit log line {line_number} is invalid JSON") from exc
            if isinstance(prior, Mapping) and prior.get("source_hash") == source_hash:
                return {
                    "status": "duplicate",
                    "appended": False,
                    "path": str(target),
                    "source_hash": source_hash,
                    "review_count": len(serializable_reviews),
                    "external_business_action": False,
                }

    event = {
        "at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "event": "action_review_import",
        "source_hash": source_hash,
        "reviewer": reviewer,
        "reviewer_claim": "self_reported",
        "reviews": serializable_reviews,
        "external_business_action": False,
    }
    line = _json(event) + "\n"
    content = existing
    if content and not content.endswith(("\n", "\r")):
        content += "\n"
    content += line

    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            newline="\n",
            dir=target.parent,
            prefix=f".{target.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            temp_path = Path(handle.name)
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_path, target)
        temp_path = None
    except OSError as exc:
        raise ClientActionError("could not atomically append audit log") from exc
    finally:
        if temp_path is not None:
            try:
                temp_path.unlink(missing_ok=True)
            except OSError:
                pass

    return {
        "status": "appended",
        "appended": True,
        "path": str(target),
        "source_hash": source_hash,
        "review_count": len(serializable_reviews),
        "external_business_action": False,
    }


__all__ = [
    "ClientActionError",
    "REVIEW_STATUSES",
    "WORKBOOK_COLUMNS",
    "WORKBOOK_HEADERS",
    "WORKBOOK_HEADER_ROW",
    "write_action_workbook",
    "load_action_reviews",
    "append_review_log",
]
