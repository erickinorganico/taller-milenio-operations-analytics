"""Rebuild the explicitly fictional two-week client walkthrough in a new folder."""
import argparse
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl
from milenio.client_input import write_client_template, load_client_input
from milenio.client_delivery import build_client, build_comparison


def build(output):
    output = Path(output)
    if output.exists():
        raise ValueError('Use una carpeta nueva')
    output.mkdir(parents=True)
    first = output / 'Entrada_semana_01.xlsx'
    second = output / 'Entrada_semana_02.xlsx'
    write_client_template(first,sample=True,as_of='2026-09-22T18:00:00Z',business_name='Taller demostrativo · datos ficticios',snapshot_id='demo-semana-01')
    wb = openpyxl.load_workbook(first)
    for row in wb['Config']:
        if row[0].value == 'as_of': row[1].value = '2026-09-29T18:00:00Z'
        if row[0].value == 'snapshot_id': row[1].value = 'demo-semana-02'
    orders = wb['Ordenes']
    for row in orders:
        if row[0].value == 'ORD-001':
            row[5].value = 'in_service'; row[6].value = '2026-09-30T18:00:00Z'; row[7].value = None
        if row[0].value == 'ORD-003':
            row[5].value = 'delivered'; row[4].value = '2026-09-26T18:00:00Z'
    payments = wb['Pagos']
    payments.append(['PAY-003','INV-001','2026-09-24T12:00:00Z',400,'transfer'])
    for table in payments.tables.values(): table.ref = f'A4:E{payments.max_row}'
    for row in wb['Inventario']:
        if row[0].value == 'PART-001': row[2].value = 8; row[3].value = 1
    wb.properties.creator = 'Milenio Analytics'; wb.properties.lastModifiedBy = 'Milenio Analytics'
    wb.save(second); wb.close()
    build_client(load_client_input(first),output / 'semana_01')
    build_client(load_client_input(second),output / 'semana_02')
    build_comparison(output / 'semana_01',output / 'semana_02',output / 'comparacion')
    return output


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',required=True)
    print(build(parser.parse_args().output))
