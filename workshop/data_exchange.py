"""Bounded file exchange through the workshop's audited write boundary."""
from __future__ import annotations

import csv
import hashlib
import io
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import ValidationError
from django.core.serializers.json import DjangoJSONEncoder
from django.db.models import Q
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from . import models as m, services as s

MAX_BYTES = 1024 * 1024
MAX_ROWS = 1000
FIELDS = {
    "customers": ("name", "phone", "email", "kind", "notes"),
    "vehicles": ("customer_name", "plate", "vin", "make", "model", "year", "odometer"),
    "parts": ("sku", "name", "unit", "cost", "sale_price", "reorder_point"),
    "suppliers": ("name", "phone", "email"),
    "orders": ("number", "vehicle_plate", "vehicle_vin", "complaint", "promised_at", "odometer"),
    "services": ("order_number", "quote_version", "description", "kind", "quantity", "unit_price", "unit_cost", "part_sku"),
    "payments": ("invoice_number", "amount", "method", "reference", "idempotency_key", "received_at"),
}
TITLES = {"customers": "Clientes", "vehicles": "Vehículos", "parts": "Refacciones",
          "suppliers": "Proveedores", "orders": "Órdenes", "services": "Servicios cotizados", "payments": "Cobros"}
MODELS = {"customers": m.Customer, "vehicles": m.Vehicle, "parts": m.Part,
          "suppliers": m.Supplier, "orders": m.WorkOrder, "services": m.QuoteLine, "payments": m.Payment}
RELATIONS = {"vehicles": ("customer",), "orders": ("vehicle",),
             "services": ("quote__work_order", "part"), "payments": ("invoice",)}
LABELS = {"id": "ID", "customer_name": "Cliente", "vehicle_plate": "Placas", "vehicle_vin": "VIN",
          "order_number": "Orden", "quote_version": "Versión de cotización", "part_sku": "SKU de refacción",
          "invoice_number": "Comprobante", "idempotency_key": "Clave única del cobro", "received_at": "Fecha del cobro"}
GUIDANCE = {
    "customers": "Tipo: individual (particular) o fleet (flotilla). Se identifican por ID o por nombre único.",
    "vehicles": "Registra primero al cliente. Usa su nombre único y al menos placa o VIN. Año y kilometraje son enteros.",
    "parts": "El SKU identifica la refacción. Usa punto decimal, sin símbolo de moneda. Las existencias se registran en Refacciones mediante entradas o ajustes.",
    "suppliers": "Se identifican por ID o por nombre único. Puedes actualizar teléfono y correo desde el archivo.",
    "orders": "Registra primero el vehículo. Las órdenes nuevas entran en Recepción. Puedes actualizar motivo, entrega prometida y kilometraje de órdenes abiertas. Las etapas se cambian desde la orden.",
    "services": "Usa el folio de la orden y la versión de cotización. Si la versión está vacía, se usa el único borrador o se crea uno durante inspección. Tipo: labor, service o part; part requiere SKU. Sólo se modifican cotizaciones en borrador.",
    "payments": "Usa un comprobante existente y una clave única por cobro. Método: cash, transfer, card u other. Repetir la misma clave y datos no duplica el pago; cambios y sobrepagos se rechazan. Fecha vacía: momento de importación.",
}


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                    separators=(",", ":"), cls=DjangoJSONEncoder).encode()).hexdigest()


def source_digest(entity):
    dependencies = {
        "customers": [m.Customer], "suppliers": [m.Supplier], "parts": [m.Part],
        "vehicles": [m.Vehicle, m.Customer], "orders": [m.WorkOrder, m.Vehicle],
        "services": [m.QuoteLine, m.Quote, m.WorkOrder, m.Part],
        "payments": [m.Payment, m.Invoice],
    }
    return digest({model.__name__: list(model.objects.order_by("pk").values()) for model in dependencies[entity]})


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, bool):
        raise ValueError("Las celdas no deben contener valores booleanos.")
    return str(value).strip()


def parse_file(upload, entity):
    if upload is None:
        return [], ["Selecciona un archivo Excel (.xlsx) o CSV."]
    if upload.size > MAX_BYTES:
        return [], ["El archivo supera 1 MB."]
    raw = upload.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        return [], ["El archivo supera 1 MB."]
    extension = Path(upload.name).suffix.casefold()
    workbook = None
    try:
        if extension == ".xlsx":
            with ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist()) > 200 or sum(item.file_size for item in archive.infolist()) > 20 * MAX_BYTES:
                    raise ValueError("El Excel supera el tamaño de lectura permitido.")
                if any("vbaproject" in item.filename.casefold() for item in archive.infolist()):
                    raise ValueError("Usa un Excel sin macros.")
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            if len(workbook.worksheets) != 1:
                raise ValueError("Usa una sola hoja por archivo; descarga la plantilla del catálogo.")
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            def excel_rows():
                for cells in sheet.iter_rows():
                    if any(cell.data_type in {"f", "e"} for cell in cells):
                        raise ValueError("El Excel contiene fórmulas o errores. Pega sólo los valores antes de subirlo.")
                    yield [_cell(cell.value) for cell in cells]
            reader = excel_rows()
        elif extension == ".csv":
            content = raw.decode("utf-8-sig", errors="strict")
            if "\x00" in content:
                raise ValueError("El CSV contiene caracteres nulos.")
            first = content.splitlines()[0] if content else ""
            delimiter = max((",", ";", "\t"), key=first.count)
            reader = csv.reader(io.StringIO(content, newline=""), delimiter=delimiter, strict=True)
        else:
            raise ValueError("Usa un archivo .xlsx o .csv. Para .xls, guarda una copia como .xlsx.")
        header = next(reader, None)
        expected = set(FIELDS[entity])
        if (header is None or len(header) != len(set(header)) or
                set(header) not in (expected, expected | {"id"})):
            raise ValueError("Los encabezados deben ser exactamente: " + ", ".join(FIELDS[entity]) + ". El ID es opcional.")
        rows = []
        for number, cells in enumerate(reader, 2):
            if number > MAX_ROWS + 1:
                raise ValueError("El archivo supera 1,000 filas de datos.")
            if len(cells) != len(header):
                # Excel drops trailing empty cells in some producers.
                if extension == ".xlsx" and len(cells) < len(header):
                    cells += [""] * (len(header) - len(cells))
                else:
                    raise ValueError(f"Fila {number}: se esperaban {len(header)} columnas y hay {len(cells)}.")
            rows.append({key: _cell(value) for key, value in zip(header, cells)})
        if not rows:
            raise ValueError("El archivo no contiene filas de datos.")
        return rows, []
    except UnicodeDecodeError:
        return [], ["El CSV debe estar codificado en UTF-8."]
    except (ValueError, csv.Error, BadZipFile, KeyError, OSError, OverflowError) as exc:
        return [], [f"Archivo inválido: {exc}"]
    except Exception:
        # Malformed XML/unsupported workbook structures are upload errors.
        return [], ["No se pudo leer el Excel. Guarda una copia como .xlsx con una sola hoja y sin fórmulas."]
    finally:
        if workbook:
            workbook.close()


def integer(raw, label, maximum=4294967295):
    if not raw:
        return None
    if not raw.isascii() or not raw.isdecimal() or len(raw) > len(str(maximum)) or int(raw) > maximum:
        raise ValidationError(f"{label}: usa un entero no negativo dentro del rango permitido.")
    return int(raw)


def decimal(raw, label, places=2, max_digits=14, positive=False):
    try:
        value = Decimal(raw)
        if not value.is_finite() or value < 0 or (positive and value == 0):
            raise InvalidOperation
        if value.as_tuple().exponent < -places or value >= Decimal(10) ** (max_digits - places):
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        raise ValidationError(f"{label}: usa un número finito {'positivo' if positive else 'no negativo'} con hasta {places} decimales.")
    return value


def timestamp(raw, label):
    if not raw:
        return None
    try:
        value = parse_datetime(raw)
    except ValueError:
        value = None
    if value is None:
        raise ValidationError(f"{label}: usa AAAA-MM-DD HH:MM, por ejemplo 2026-10-01 09:30.")
    return timezone.make_aware(value) if timezone.is_naive(value) else value


def unique(qs, label, required=True):
    matches = list(qs[:2])
    if len(matches) > 1 or (not matches and required):
        raise ValidationError(f"{label}: debe coincidir con un único registro existente.")
    return matches[0] if matches else None


def snapshot(obj):
    return {field.name: str(getattr(obj, field.attname)) if getattr(obj, field.attname) is not None else None
            for field in obj._meta.fields}


def _target(entity, row):
    model = MODELS[entity]
    if row.get("id"):
        pk = integer(row["id"], "ID", 9223372036854775807)
        return unique(model.objects.filter(pk=pk), "ID")
    if entity in {"customers", "suppliers"}:
        return unique(model.objects.filter(name__iexact=row["name"]), "Nombre", False)
    if entity == "parts":
        return unique(model.objects.filter(sku__iexact=row["sku"]), "SKU", False)
    if entity == "vehicles":
        query = Q(pk__in=[])
        if row["plate"]:
            query |= Q(plate__iexact=row["plate"])
        if row["vin"]:
            query |= Q(vin__iexact=row["vin"])
        return unique(model.objects.filter(query), "Placa o VIN", False)
    if entity == "orders":
        return unique(model.objects.filter(number__iexact=row["number"]), "Folio de orden", False)
    if entity == "payments":
        return unique(model.objects.filter(idempotency_key=row["idempotency_key"]), "Clave del cobro", False)
    return None


def _catalog_values(entity, row):
    if entity in {"customers", "suppliers"}:
        return {field: row[field] for field in FIELDS[entity]}
    if entity == "parts":
        return {"sku": row["sku"].upper(), "name": row["name"], "unit": row["unit"],
                "cost": decimal(row["cost"], "Costo"), "sale_price": decimal(row["sale_price"], "Precio"),
                "reorder_point": decimal(row["reorder_point"], "Reposición", 3, 12)}
    if not row["plate"] and not row["vin"]:
        raise ValidationError("Se requiere placa o VIN.")
    year = integer(row["year"], "Año", 65535)
    if year is not None and not 1886 <= year <= timezone.localdate().year + 1:
        raise ValidationError("Indica un año de vehículo válido.")
    return {"customer": unique(m.Customer.objects.filter(name__iexact=row["customer_name"]), "Cliente"),
            "plate": row["plate"].upper(), "vin": row["vin"].upper() or None,
            "make": row["make"], "model": row["model"], "year": year,
            "odometer": integer(row["odometer"], "Kilometraje")}


def _check_unique(entity, obj):
    checks = {"customers": ["name"], "suppliers": ["name"], "parts": ["sku"], "vehicles": ["plate", "vin"], "orders": ["number"]}
    for field in checks.get(entity, []):
        value = getattr(obj, field)
        if value and type(obj).objects.filter(**{field + "__iexact": value}).exclude(pk=obj.pk).exists():
            raise ValidationError(f"{field}: ya existe en otro registro.")


def _save_catalog(actor, entity, obj, row):
    before = snapshot(obj) if obj else {}
    obj = obj or MODELS[entity]()
    for field, value in _catalog_values(entity, row).items():
        setattr(obj, field, value)
    _check_unique(entity, obj)
    obj.full_clean()
    after = snapshot(obj)
    if before == after:
        return obj, before, "unchanged"
    obj.save()
    m.AuditEvent.objects.create(actor=actor, entity_type=type(obj).__name__, entity_id=str(obj.pk),
                               action="file_updated" if before else "file_created", before=before, after=snapshot(obj))
    return obj, before, "updated" if before else "created"


def _order(actor, obj, row):
    query = Q(pk__in=[])
    if row["vehicle_plate"]:
        query |= Q(plate__iexact=row["vehicle_plate"])
    if row["vehicle_vin"]:
        query |= Q(vin__iexact=row["vehicle_vin"])
    vehicle = unique(m.Vehicle.objects.filter(query), "Vehículo")
    if ((row["vehicle_plate"] and vehicle.plate.casefold() != row["vehicle_plate"].casefold()) or
            (row["vehicle_vin"] and (vehicle.vin or "").casefold() != row["vehicle_vin"].casefold())):
        raise ValidationError("La placa y el VIN deben corresponder al mismo vehículo.")
    values = {"number": row["number"], "vehicle": vehicle, "complaint": row["complaint"],
              "promised_at": timestamp(row["promised_at"], "Entrega prometida"), "odometer": integer(row["odometer"], "Kilometraje")}
    candidate = obj or m.WorkOrder()
    before = snapshot(candidate) if obj else {}
    for field, value in values.items():
        setattr(candidate, field, value)
    _check_unique("orders", candidate)
    candidate.full_clean()
    if before == snapshot(candidate):
        return candidate, before, "unchanged"
    if not obj:
        return s.create_work_order(actor=actor, **values), before, "created"
    if candidate.status in {"delivered", "cancelled"}:
        raise ValidationError("Las órdenes entregadas o canceladas se consultan; no se actualizan desde un archivo.")
    if before["vehicle"] != str(vehicle.pk) and candidate.status != "intake":
        raise ValidationError("Sólo se cambia el vehículo de una orden en Recepción.")
    candidate.version += 1
    candidate.save()
    m.AuditEvent.objects.create(actor=actor, entity_type="WorkOrder", entity_id=str(candidate.pk),
                               action="file_updated", before=before, after=snapshot(candidate))
    return candidate, before, "updated"


def _service(actor, obj, row):
    order = unique(m.WorkOrder.objects.filter(number__iexact=row["order_number"]), "Orden")
    version = integer(row["quote_version"], "Versión", 4294967295)
    if version is not None:
        quote = unique(m.Quote.objects.filter(work_order=order, version=version), "Cotización")
    else:
        quote = unique(m.Quote.objects.filter(work_order=order, status="draft"), "Borrador", False)
        if quote is None:
            quote = s.create_quote(actor=actor, work_order=order)
    part = unique(m.Part.objects.filter(sku__iexact=row["part_sku"]), "SKU") if row["part_sku"] else None
    if row["kind"] not in {"labor", "service", "part"} or (row["kind"] == "part") != bool(part):
        raise ValidationError("Tipo: labor, service o part. Sólo part requiere un SKU de refacción.")
    if obj is None:
        obj = unique(m.QuoteLine.objects.filter(quote=quote, description=row["description"], kind=row["kind"], part=part), "Servicio; usa el ID si hay descripciones repetidas", False)
    values = {"description": row["description"], "kind": row["kind"],
              "quantity": decimal(row["quantity"], "Cantidad", 3, 12, True),
              "unit_price": decimal(row["unit_price"], "Precio"),
              "unit_cost": decimal(row["unit_cost"], "Costo") if row["unit_cost"] else None, "part": part}
    before = snapshot(obj) if obj else {}
    candidate = obj or m.QuoteLine(quote=quote)
    if candidate.quote_id != quote.pk:
        raise ValidationError("El ID del servicio corresponde a otra cotización.")
    for field, value in values.items():
        setattr(candidate, field, value)
    candidate.full_clean()
    if before == snapshot(candidate):
        return candidate, before, "unchanged"
    if quote.status != "draft":
        raise ValidationError("Sólo se modifican servicios de cotizaciones en borrador.")
    if not obj:
        return s.add_quote_line(actor=actor, quote=quote, **values), before, "created"
    candidate.save()
    m.AuditEvent.objects.create(actor=actor, entity_type="QuoteLine", entity_id=str(candidate.pk),
                               action="file_updated", before=before, after=snapshot(candidate))
    return candidate, before, "updated"


def _payment(actor, obj, row):
    invoice = unique(m.Invoice.objects.filter(number__iexact=row["invoice_number"]), "Comprobante")
    if row["method"] not in {"cash", "transfer", "card", "other"}:
        raise ValidationError("Método: cash, transfer, card u other.")
    candidate = m.Payment(invoice=invoice, amount=decimal(row["amount"], "Importe", positive=True),
                          method=row["method"], reference=row["reference"], idempotency_key=row["idempotency_key"],
                          received_at=timestamp(row["received_at"], "Fecha del cobro") or timezone.now(), created_by=actor)
    candidate.full_clean(validate_unique=False)
    if obj and obj.idempotency_key != candidate.idempotency_key:
        raise ValidationError("El ID y la clave única del cobro no corresponden.")
    if not obj and m.Payment.objects.filter(idempotency_key=candidate.idempotency_key).exists():
        raise ValidationError("La clave única corresponde a otro cobro.")
    result = s.record_payment(actor=actor, invoice=invoice, amount=candidate.amount, method=candidate.method,
                              reference=candidate.reference, idempotency_key=candidate.idempotency_key,
                              received_at=timestamp(row["received_at"], "Fecha del cobro"))
    return result, snapshot(obj) if obj else {}, "unchanged" if obj else "created"


def apply_rows(actor, entity, rows, mode):
    """Caller owns atomic transaction; preview caller rolls back all writes."""
    results, seen = [], set()
    for number, row in enumerate(rows, 2):
        try:
            obj = _target(entity, row)
            if obj and mode == "create":
                raise ValidationError("El registro ya existe. Selecciona Agregar y actualizar para revisarlo.")
            if entity == "orders":
                obj, before, action = _order(actor, obj, row)
            elif entity == "services":
                obj, before, action = _service(actor, obj, row)
                if before and mode == "create":
                    raise ValidationError("El servicio ya existe. Selecciona Agregar y actualizar.")
            elif entity == "payments":
                obj, before, action = _payment(actor, obj, row)
            else:
                obj, before, action = _save_catalog(actor, entity, obj, row)
            if obj.pk in seen:
                raise ValidationError("El registro se repite en el archivo.")
            seen.add(obj.pk)
            after = snapshot(obj)
            changes = [{"field": field, "before": before.get(field), "after": value}
                       for field, value in after.items() if field not in {"id", "updated_at", "version", "created_at"}
                       and before.get(field) != value]
            results.append({"row": row, "action": action, "changes": changes, "id": obj.pk})
        except ValidationError as exc:
            detail = "; ".join(f"{field}: {', '.join(values)}" for field, values in exc.message_dict.items()) if hasattr(exc, "message_dict") else "; ".join(exc.messages)
            raise ValidationError(f"Fila {number}: {detail}")
    return results


def export_rows(entity, ids=None, objects=None):
    if objects is None:
        objects = MODELS[entity].objects.order_by("pk")
        if entity in RELATIONS:
            objects = objects.select_related(*RELATIONS[entity])
        if ids is not None:
            objects = objects.filter(pk__in=ids)
    for obj in objects:
        if entity == "vehicles":
            values = {field: getattr(obj, field) for field in FIELDS[entity] if field != "customer_name"}
            values["customer_name"] = obj.customer.name
        elif entity == "orders":
            values = {field: getattr(obj, field) for field in FIELDS[entity] if not field.startswith("vehicle_")}
            values.update(vehicle_plate=obj.vehicle.plate, vehicle_vin=obj.vehicle.vin)
        elif entity == "services":
            values = {field: getattr(obj, field) for field in FIELDS[entity] if field not in {"order_number", "quote_version", "part_sku"}}
            values.update(order_number=obj.quote.work_order.number, quote_version=obj.quote.version, part_sku=obj.part.sku if obj.part_id else "")
        elif entity == "payments":
            values = {field: getattr(obj, field) for field in FIELDS[entity] if field != "invoice_number"}
            values["invoice_number"] = obj.invoice.number
        else:
            values = {field: getattr(obj, field) for field in FIELDS[entity]}
        for field, value in values.items():
            if isinstance(value, datetime):
                values[field] = timezone.localtime(value).isoformat()
        yield {"id": obj.pk, **values}


def file_content(entity, format, *, template=False):
    fields = ("id",) + FIELDS[entity] if not template else FIELDS[entity]
    rows = [] if template else export_rows(entity)
    if format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(fields)
        for row in rows:
            cells = [_cell(row.get(field)) for field in fields]
            writer.writerow(["'" + cell if cell.lstrip().startswith(("=", "+", "-", "@")) else cell for cell in cells])
        return ("\ufeff" + output.getvalue()).encode("utf-8"), "text/csv; charset=utf-8"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TITLES[entity]
    sheet.append(fields)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244C3B")
    for row in rows:
        sheet.append([_cell(row.get(field)) for field in fields])
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"  # A literal starting with '=' is never an Excel formula.
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
