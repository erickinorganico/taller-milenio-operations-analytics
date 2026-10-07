"""File updates must preserve domain invariants and match the reviewed source."""
import csv
import io
from datetime import timedelta
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from workshop import models as m
from workshop.data_exchange import FIELDS, MAX_ROWS, file_content


def upload(entity, rows, xlsx=False, delimiter=","):
    fields = (("id",) + FIELDS[entity]) if any("id" in row for row in rows) else FIELDS[entity]
    if xlsx:
        book = Workbook()
        sheet = book.active
        sheet.append(fields)
        for row in rows:
            sheet.append([row.get(key, "") for key in fields])
        stream = io.BytesIO()
        book.save(stream)
        return SimpleUploadedFile("datos.xlsx", stream.getvalue())
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, delimiter=delimiter)
    writer.writerow(fields)
    for row in rows:
        writer.writerow([row.get(key, "") for key in fields])
    return SimpleUploadedFile("datos.csv", stream.getvalue().encode("utf-8"))


class DataExchangeTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.manager = User.objects.create_superuser("exchange-manager", password="TestOnlyPassword!2026")
        cls.customer = m.Customer.objects.create(name="Cliente original", phone="001234", kind="individual")
        cls.vehicle = m.Vehicle.objects.create(customer=cls.customer, plate="ABC123", make="Marca", model="Modelo")

    def setUp(self):
        self.client.force_login(self.manager)

    def preview(self, entity, rows, **kwargs):
        return self.client.post("/import/", {"entity": entity, "mode": "upsert", "stage": "preview", "file": upload(entity, rows, **kwargs)})

    def confirm(self, entity, response):
        return self.client.post("/import/", {"entity": entity, "mode": "upsert", "stage": "commit", "receipt": response.context["receipt"]})

    def customer_row(self, **values):
        return {"id": self.customer.pk, "name": "Cliente original", "phone": "009999", "email": "", "kind": "individual", "notes": "", **values}

    def order(self, status="inspection"):
        return m.WorkOrder.objects.create(number="OT-1", vehicle=self.vehicle, complaint="Servicio", status=status)

    def test_excel_update_and_insert_are_previewed_then_visible_with_new_snapshot(self):
        initial_audits = m.AuditEvent.objects.count()
        response = self.preview("customers", [self.customer_row(name="Cliente actualizado"),
            self.customer_row(id="", name="Cliente nuevo", phone="00001")], xlsx=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["summary"], {"created": 1, "updated": 1, "unchanged": 0})
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.name, "Cliente original")
        self.assertEqual(m.Customer.objects.count(), 1)
        self.assertEqual(m.AuditEvent.objects.count(), initial_audits)
        self.assertEqual(m.AnalyticsSnapshot.objects.count(), 0)
        self.assertEqual(self.confirm("customers", response).status_code, 302)
        self.customer.refresh_from_db()
        self.assertEqual((self.customer.name, self.customer.phone), ("Cliente actualizado", "009999"))
        self.assertEqual(m.Customer.objects.count(), 2)
        self.assertEqual(m.AnalyticsSnapshot.objects.get().source_counts["Customer"], 2)
        self.assertContains(self.client.get("/data/customers/?q=actualizado"), "Cliente actualizado")
        self.assertContains(self.client.get("/data/"), "Base de datos")

    def test_stale_existing_record_and_cross_user_receipt_cannot_overwrite(self):
        response = self.preview("customers", [self.customer_row()])
        m.Customer.objects.filter(pk=self.customer.pk).update(phone="Cambio del equipo")
        self.assertEqual(self.confirm("customers", response).status_code, 409)
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.phone, "Cambio del equipo")
        self.assertEqual(m.AuditEvent.objects.count(), 0)
        response = self.preview("customers", [self.customer_row()])
        second = User.objects.create_superuser("second-manager", password="TestOnlyPassword!2026")
        self.client.force_login(second)
        self.assertEqual(self.confirm("customers", response).status_code, 400)

    def test_csv_semicolon_and_exported_excel_round_trip_keep_text_and_ids(self):
        response = self.preview("suppliers", [{"name": "Proveedor Ñ", "phone": "00010", "email": ""}], delimiter=";")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.confirm("suppliers", response).status_code, 302)
        exported = self.client.get("/data/export/customers/?format=xlsx")
        self.assertEqual(exported.status_code, 200)
        book = load_workbook(io.BytesIO(exported.content))
        sheet = book.active
        self.assertEqual(sheet["A2"].value, str(self.customer.pk))
        self.assertEqual(sheet["C2"].value, "001234")
        response = self.client.post("/import/", {"entity": "customers", "mode": "upsert", "stage": "preview",
            "file": SimpleUploadedFile("actuales.xlsx", exported.content)})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["summary"]["unchanged"], 1)

    def test_orders_enter_reception_and_service_updates_respect_quote_authorization(self):
        response = self.preview("orders", [{"number": "OT-NUEVA", "vehicle_plate": "ABC123", "vehicle_vin": "",
            "complaint": "Frenos", "promised_at": "2026-10-10 12:00", "odometer": "1234"}])
        self.assertEqual(response.status_code, 200)
        self.assertFalse(m.WorkOrder.objects.exists())
        self.confirm("orders", response)
        self.assertEqual(m.WorkOrder.objects.get().status, "intake")
        order = self.order()
        row = {"order_number": order.number, "quote_version": "", "description": "Alineación", "kind": "labor",
               "quantity": "1", "unit_price": "350.00", "unit_cost": "100.00", "part_sku": ""}
        response = self.preview("services", [row])
        self.assertEqual(response.status_code, 200)
        self.assertFalse(m.Quote.objects.exists())
        self.assertEqual(self.confirm("services", response).status_code, 302)
        line = m.QuoteLine.objects.get()
        self.assertEqual(line.unit_price, Decimal("350.00"))
        response = self.preview("services", [{**row, "id": line.pk, "quote_version": "1", "unit_price": "400.00"}])
        self.assertEqual(response.status_code, 200)
        self.confirm("services", response)
        line.refresh_from_db()
        self.assertEqual(line.unit_price, Decimal("400.00"))
        self.assertContains(self.client.get("/data/services/?q=OT-1"), "Alineación")
        self.assertEqual(self.client.get("/data/services/?view=source").status_code, 200)
        m.Quote.objects.filter(pk=line.quote_id).update(status="approved")
        invalid = self.preview("services", [{**row, "id": line.pk, "quote_version": "1", "unit_price": "500.00"}])
        self.assertEqual(invalid.status_code, 400)
        line.refresh_from_db()
        self.assertEqual(line.unit_price, Decimal("400.00"))

    def test_payments_replay_and_total_file_overpayment_are_atomic(self):
        order = self.order("delivered")
        invoice = m.Invoice.objects.create(work_order=order, number="COM-1", subtotal=100, tax=0, total=100,
                                          due_at=timezone.now() + timedelta(days=2))
        row = {"invoice_number": invoice.number, "amount": "60.00", "method": "transfer",
               "reference": "REC-1", "idempotency_key": "COBRO-1", "received_at": "2026-10-01 10:00"}
        invalid = self.preview("payments", [row, {**row, "idempotency_key": "COBRO-2"}])
        self.assertEqual(invalid.status_code, 400)
        self.assertFalse(m.Payment.objects.exists())
        self.assertFalse(m.AuditEvent.objects.exists())
        response = self.preview("payments", [row])
        self.assertEqual(response.status_code, 200)
        self.assertFalse(m.Payment.objects.exists())
        self.assertEqual(self.confirm("payments", response).status_code, 302)
        response = self.preview("payments", [row])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["summary"]["unchanged"], 1)
        self.confirm("payments", response)
        self.assertEqual(m.Payment.objects.count(), 1)
        self.assertEqual(self.preview("payments", [{**row, "amount": "50.00"}]).status_code, 400)
        exported = self.client.get("/data/export/payments/?format=xlsx")
        self.assertEqual(exported.status_code, 200)
        self.assertContains(self.client.get("/data/payments/"), "COBRO-1")

    def test_invalid_workbook_limits_and_id_collisions_leave_everything_unchanged(self):
        formula = self.preview("customers", [self.customer_row(name="=1+1")], xlsx=True)
        self.assertEqual(formula.status_code, 400)
        collision = self.preview("customers", [self.customer_row(), self.customer_row(phone="00222")])
        self.assertEqual(collision.status_code, 400)
        invalid = self.client.post("/import/", {"entity": "customers", "mode": "upsert", "stage": "preview",
            "file": SimpleUploadedFile("roto.xlsx", b"not an xlsx")})
        self.assertEqual(invalid.status_code, 400)
        self.customer.refresh_from_db()
        self.assertEqual(self.customer.phone, "001234")
        self.assertFalse(m.AuditEvent.objects.exists())
        many = [self.customer_row(id="", name=f"Cliente {i}") for i in range(MAX_ROWS + 1)]
        self.assertEqual(self.preview("customers", many, xlsx=True).status_code, 400)

    def test_exports_escape_formulas_and_all_read_write_roles_remain_enforced(self):
        self.customer.notes = '=HYPERLINK("https://example.invalid","click")'
        self.customer.save()
        content, _ = file_content("customers", "xlsx")
        book = load_workbook(io.BytesIO(content))
        self.assertEqual(book.active["F2"].data_type, "s")
        content, _ = file_content("customers", "csv")
        self.assertIn("'=HYPERLINK", content.decode("utf-8-sig"))
        viewer = User.objects.create_user("viewer-exchange")
        viewer.groups.add(Group.objects.create(name="viewer"))
        self.client.force_login(viewer)
        self.assertEqual(self.client.get("/data/").status_code, 200)
        self.assertEqual(self.client.get("/data/export/orders/").status_code, 200)
        self.assertEqual(self.client.get("/import/").status_code, 403)
        technician = User.objects.create_user("technician-exchange")
        technician.groups.add(Group.objects.create(name="technician"))
        self.client.force_login(technician)
        self.assertEqual(self.client.get("/data/").status_code, 403)
        self.assertEqual(self.client.get("/data/export/customers/").status_code, 403)

    def test_expired_preview_and_mismatched_plate_vin_are_rejected(self):
        response = self.preview("customers", [self.customer_row()])
        with patch("django.core.signing.time.time", return_value=timezone.now().timestamp() + 1000):
            self.assertEqual(self.confirm("customers", response).status_code, 400)
        m.Vehicle.objects.create(customer=self.customer, plate="XYZ", vin="VIN2", make="Marca", model="Otro")
        row = {"number": "OT-X", "vehicle_plate": "ABC123", "vehicle_vin": "VIN2", "complaint": "X", "promised_at": "", "odometer": ""}
        self.assertEqual(self.preview("orders", [row]).status_code, 400)
        self.assertFalse(m.WorkOrder.objects.exists())
