import io
import tempfile
from datetime import timedelta
from pathlib import Path

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook
from PIL import Image

from workshop import models as m
from workshop.document_capture import process_one, suggest_fields


TEXT = "Orden: FOTO-001\nCliente: Cliente de prueba\nPlacas: ABC123\nKilometraje: 0\nSolicitud: Ruido\nServicio: Frenos\nServicio: Aceite"


def photo(name="orden.png"):
    stream = io.BytesIO()
    Image.new("RGB", (120, 120), "white").save(stream, "PNG")
    return SimpleUploadedFile(name, stream.getvalue(), content_type="image/png")


class DocumentCaptureTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.settings = override_settings(MEDIA_ROOT=self.temp.name)
        self.settings.enable()
        self.addCleanup(self.settings.disable)
        self.manager = User.objects.create_superuser("capture-manager")
        self.client.force_login(self.manager)

    def upload(self, name="orden.png"):
        response = self.client.post("/capture/", {"title":"Hoja de prueba", "kind":"order", "photo":photo(name)})
        self.assertEqual(response.status_code, 302)
        return m.DocumentCapture.objects.get()

    def review(self, capture, **values):
        return self.client.post(f"/capture/{capture.pk}/", {"version":capture.version, "order_number":"FOTO-001",
            "customer_name":"=SUM(1,2)", "vehicle_plate":"ABC123", "odometer":"0", "complaint":"Ruido",
            "services":"Frenos\n=malicious-formula", "checked":"on", **values})

    def test_photo_review_excel_and_csv_do_not_change_operational_records(self):
        counts = (m.WorkOrder.objects.count(), m.Payment.objects.count(), m.QuoteLine.objects.count())
        capture = self.upload()
        self.assertEqual(Path(capture.photo.name).suffix, ".jpg")
        self.assertEqual(self.client.get(f"/capture/{capture.pk}/export/xlsx/").status_code, 404)
        self.assertTrue(process_one(reader=lambda path: TEXT))
        capture.refresh_from_db()
        self.assertEqual(capture.status, "review")
        self.assertEqual(capture.data["services"], "Frenos\nAceite")
        self.assertEqual(self.review(capture).status_code, 302)
        capture.refresh_from_db()
        self.assertEqual(capture.status, "confirmed")
        self.assertEqual(capture.data["odometer"], "0")
        exported = self.client.get(f"/capture/{capture.pk}/export/xlsx/")
        book = load_workbook(io.BytesIO(exported.content))
        self.assertEqual(book.sheetnames, ["Orden revisada", "Servicios revisados"])
        self.assertEqual(book.worksheets[0]["B3"].value, "=SUM(1,2)")
        self.assertEqual(book.worksheets[0]["B3"].data_type, "s")
        self.assertEqual(book.worksheets[1]["B3"].data_type, "s")
        csv = self.client.get(f"/capture/{capture.pk}/export/csv/")
        self.assertIn("'=SUM(1,2)", csv.content.decode("utf-8-sig"))
        self.assertEqual(counts, (m.WorkOrder.objects.count(), m.Payment.objects.count(), m.QuoteLine.objects.count()))
        self.assertEqual(m.AuditEvent.objects.filter(entity_type="DocumentCapture", action="confirmed").count(), 1)

    def test_structured_service_export_keeps_numbers_sku_and_unknowns(self):
        import csv
        from workshop.document_capture import export_content
        capture = self.upload()
        capture.reviewed_by = self.manager
        capture.reviewed_at = timezone.now()
        capture.data = {'order_number':'SYNTHETIC','services':'Texto'}
        capture.service_items = [{'description':'Servicio','quantity':2,'unit_price':'123.45','unit_cost':'67.89','sku':'=SKU'},
                                 {'description':'Pendiente','quantity':None,'unit_price':None,'unit_cost':None,'sku':''}]
        rows = list(csv.DictReader(io.StringIO(export_content(capture,'csv')[0].decode('utf-8-sig'))))
        self.assertEqual(rows[0]['service_quantity'],'2')
        self.assertEqual(rows[0]['service_unit_price'],'123.45')
        self.assertEqual(rows[0]['service_unit_cost'],'67.89')
        self.assertEqual(rows[0]['service_sku'],"'=SKU")
        self.assertEqual(rows[1]['service_quantity'],'')
        book = load_workbook(io.BytesIO(export_content(capture,'xlsx')[0]))
        self.assertEqual(book['Servicios revisados']['H2'].value,'=SKU')
        self.assertEqual(book['Servicios revisados']['H2'].data_type,'s')

    def test_duplicates_are_reused_and_roles_protect_photos_exports_and_review(self):
        capture = self.upload()
        self.upload()
        self.assertEqual(m.DocumentCapture.objects.count(), 1)
        viewer = User.objects.create_user("capture-viewer")
        viewer.groups.add(Group.objects.create(name="viewer"))
        self.client.force_login(viewer)
        for suffix in ["", "status/", "photo/", "export/xlsx/"]:
            self.assertEqual(self.client.get(f"/capture/{capture.pk}/{suffix}").status_code, 403)
        self.assertEqual(self.review(capture).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(f"/capture/{capture.pk}/photo/").status_code, 302)

    def test_failure_empty_text_and_interrupted_read_can_be_reviewed_manually(self):
        capture = self.upload()
        def unavailable(path):
            raise RuntimeError("OCR unavailable")
        process_one(reader=unavailable)
        capture.refresh_from_db()
        self.assertEqual(capture.status, "review")
        self.assertIn("manualmente", capture.reading_note)
        capture.status = "reading"
        capture.reading_started_at = timezone.now()-timedelta(minutes=4)
        capture.reading_token = "expired"
        capture.save()
        self.assertFalse(process_one(reader=lambda path: ""))
        capture.refresh_from_db()
        self.assertEqual(capture.status, "review")
        self.assertEqual(capture.reading_token, "")
        self.assertIn("interrumpió", capture.reading_note)

    def test_paused_queue_has_manual_fallback_and_late_ocr_cannot_replace_it(self):
        capture = self.upload()
        policy = m.AutomationPolicy.objects.create(pk=1, paused=True)
        self.assertFalse(process_one(reader=lambda path: TEXT))
        self.assertEqual(self.client.post(f"/capture/{capture.pk}/manual/", {"version":capture.version}).status_code, 302)
        capture.refresh_from_db()
        self.assertEqual(capture.status, "review")
        capture.status = "pending"
        capture.save()
        policy.paused = False
        policy.save()
        def late_reader(path):
            current = m.DocumentCapture.objects.get(pk=capture.pk)
            self.client.post(f"/capture/{capture.pk}/manual/", {"version":current.version})
            return TEXT
        process_one(reader=late_reader)
        capture.refresh_from_db()
        self.assertEqual(capture.raw_text, "")
        self.assertEqual(capture.status, "review")
        self.assertIn("Captura manual", capture.reading_note)

    def test_stale_or_unfinished_review_cannot_overwrite_and_confirmation_is_required(self):
        capture = self.upload()
        self.assertEqual(self.review(capture).status_code, 409)
        process_one(reader=lambda path: TEXT)
        capture.refresh_from_db()
        self.assertEqual(self.review(capture, checked="").status_code, 200)
        self.assertEqual(self.review(capture, version=capture.version-1).status_code, 409)
        self.assertEqual(self.review(capture).status_code, 302)
        self.assertEqual(self.review(capture, complaint="stale").status_code, 409)
        self.assertEqual(self.client.post(f"/capture/{capture.pk}/retry/", {"version":"bad"}).status_code, 400)

    def test_invalid_image_and_unlabelled_text_do_not_create_false_data(self):
        invalid = SimpleUploadedFile("photo.jpg", b"not an image", content_type="image/jpeg")
        self.assertEqual(self.client.post("/capture/", {"title":"Inválida", "kind":"order", "photo":invalid}).status_code, 200)
        self.assertFalse(m.DocumentCapture.objects.exists())
        self.assertEqual(self.client.post("/capture/", {"title":"Extensión inválida", "kind":"order", "photo":photo("fake.html")}).status_code, 200)
        self.assertFalse(m.DocumentCapture.objects.exists())
        self.assertTrue(all(not value for value in suggest_fields("ruido frenos 85000 sin etiquetas").values()))
        self.assertEqual(suggest_fields(TEXT)["odometer"], "0")

    def test_example_is_demo_only_and_actual_photo_is_private(self):
        with override_settings(MILENIO_MODE="live"):
            self.assertEqual(self.client.post("/capture/example/").status_code, 404)
        with override_settings(MILENIO_MODE="demo"):
            self.assertEqual(self.client.post("/capture/example/").status_code, 302)
            self.assertEqual(self.client.post("/capture/example/").status_code, 302)
        capture = m.DocumentCapture.objects.get()
        response = self.client.get(f"/capture/{capture.pk}/photo/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "image/jpeg")
        response.close()

    def test_attention_filters_exclude_closed_orders_and_order_has_next_step(self):
        customer = m.Customer.objects.create(name="Cliente")
        vehicle = m.Vehicle.objects.create(customer=customer, plate="ABC123")
        open_order = m.WorkOrder.objects.create(number="SIN-RESP", vehicle=vehicle, complaint="Inspección")
        m.WorkOrder.objects.create(number="CERRADA", vehicle=vehicle, complaint="Completada", status="delivered")
        response = self.client.get("/orders/?attention=unassigned")
        self.assertContains(response, "SIN-RESP")
        self.assertNotContains(response, "CERRADA")
        self.assertContains(self.client.get(f"/orders/{open_order.pk}/"), "Iniciar la inspección")
        self.assertContains(self.client.get("/today/"), "1 órdenes sin responsable")
