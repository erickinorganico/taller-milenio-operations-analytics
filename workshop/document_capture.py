"""Local OCR produces suggestions; only a human review permits an export."""
import csv
from decimal import Decimal, InvalidOperation
import io
import json
import os
import re
import subprocess
import tempfile
import uuid
from datetime import timedelta
from pathlib import Path

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import F
from django.utils import timezone
from PIL import Image, ImageOps
from openpyxl import Workbook

from .models import AuditEvent, DocumentCapture, AutomationPolicy


LABELS = {
    "order_number": "Folio de orden", "customer_name": "Cliente", "phone": "Teléfono",
    "vehicle_plate": "Placas", "vehicle_description": "Vehículo", "odometer": "Kilometraje",
    "complaint": "Solicitud / observaciones", "services": "Servicios (uno por renglón)",
    "vehicle_vin": "VIN", "vehicle_make": "Marca", "vehicle_model": "Modelo", "vehicle_year": "Año",
}
PATTERNS = {
    "order_number": r"(?:orden(?: de trabajo)?|folio)(?:\s*(?:no\.?|n[uú]mero|#))?\s*[:#-]\s*(.+)",
    "customer_name": r"cliente\s*:\s*(.+)", "phone": r"tel[eé]fono\s*:\s*(.+)",
    "vehicle_plate": r"placas?\s*:\s*(.+)", "vehicle_description": r"veh[ií]culo\s*:\s*(.+)",
    "odometer": r"(?:kilometraje|km)\s*:\s*([\d ,]+)",
    "complaint": r"(?:solicitud|observaciones|falla|motivo)\s*:\s*(.+)",
    "vehicle_vin": r"(?:vin|serie)\s*:\s*(.+)", "vehicle_make": r"marca\s*:\s*(.+)",
    "vehicle_model": r"modelo\s*:\s*(.+)", "vehicle_year": r"a[ñn]o\s*:\s*(\d{4})",
}


def suggest_fields(text):
    """Use explicit labels, leaving uncertain/unlabelled information blank."""
    data = {key: "" for key in LABELS}
    for key, pattern in PATTERNS.items():
        found = re.search(r"^\s*" + pattern + r"\s*$", text, re.I | re.M)
        if found:
            data[key] = found[1].strip()[:160 if key != "complaint" else 4000]
    mileage = data["odometer"].replace(",", "").replace(" ", "")
    data["odometer"] = mileage if mileage.isascii() and mileage.isdigit() else ""
    data["services"] = "\n".join(match[1].strip() for match in re.finditer(r"^\s*servicio\s*:\s*(.+)$", text, re.I | re.M))[:8000]
    return data


def read_photo(path,cancel_event=None):
    if os.name != "nt":
        raise RuntimeError("La lectura local del piloto requiere Windows; puedes capturar los campos manualmente.")
    powershell = Path(os.environ.get("SystemRoot", r"C:\Windows")) / "System32/WindowsPowerShell/v1.0/powershell.exe"
    script = Path(settings.BASE_DIR) / "scripts/read_document_photo.ps1"
    # Normalize orientation and dimensions; OCR never receives image metadata.
    with tempfile.TemporaryDirectory(prefix="milenio-ocr-") as folder:
        normalized = Path(folder) / "document.png"
        with Image.open(path) as original:
            image = ImageOps.exif_transpose(original).convert("RGB")
            image.thumbnail((2400, 2400))
            image.save(normalized, format="PNG")
        command = [str(powershell), "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                   "-File", str(script), "-ImagePath", str(normalized)]
        if cancel_event is not None:
            from .capture_vision import _run_cancelable
            result = _run_cancelable(command,"",cancel_event,timeout=25)
        else:
            result = subprocess.run(command,capture_output=True,timeout=25,
                encoding="utf-8-sig",errors="replace",creationflags=subprocess.CREATE_NO_WINDOW)
        if result.returncode:
            raise RuntimeError("No fue posible leer la foto con el OCR local. Puedes completar los campos manualmente o volver a intentar con una foto más clara.")
        output = json.loads(result.stdout)
        text = output.get("text", "")
        if not isinstance(text, str) or len(text) > 50000:
            raise RuntimeError("La lectura excedió el tamaño permitido; usa una foto de una sola hoja.")
        return text


def process_one(reader=None,cancel_event=None):
    """Claim one persisted draft and read outside the database transaction."""
    if AutomationPolicy.objects.filter(pk=1, paused=True).exists():
        return False
    with transaction.atomic():
        DocumentCapture.objects.filter(status="reading", reading_started_at__lt=timezone.now()-timedelta(minutes=3)).update(
            status="review", reading_token="", reading_note="La lectura se interrumpió. Reintenta o captura manualmente.", version=F("version")+1)
        capture = DocumentCapture.objects.select_for_update().filter(status="pending").order_by("created_at", "pk").first()
        if not capture:
            return False
        token = uuid.uuid4().hex
        capture.status, capture.reading_token = "reading", token
        capture.reading_started_at = timezone.now()
        capture.save(update_fields=["status", "reading_token", "reading_started_at"])
    try:
        if reader or capture.provider == "local":
            text = reader(capture.photo.path) if reader else read_photo(capture.photo.path,cancel_event=cancel_event)
            data = suggest_fields(text)
            items = []
            extraction = {"provider":"local","model":"windows-ocr","completion_observed":bool(text.strip())}
            note = "Lectura local terminada. Compara cada campo con la foto; la escritura a mano puede requerir captura manual." if text.strip() else "No se detectó texto. Completa los campos manualmente o sube una foto más clara."
        else:
            from .capture_vision import extract_photo
            output = extract_photo(capture.photo.path,provider=capture.provider,model=capture.model_name,cancel_event=cancel_event)
            data, items = output["data"], output["service_items"]
            data["services"] = "\n".join(item["description"] for item in items)
            text = ""
            extraction = {**output["receipt"],"warnings":output["warnings"]}
            note = "El modelo leyó la imagen. Revisa las coincidencias y los campos propuestos antes de aplicar al taller."
    except Exception as error:
        text, data = "", suggest_fields("")
        items = []
        extraction = {"provider":capture.provider,"model":capture.model_name,"completion_observed":False,"error_type":type(error).__name__}
        note = "No se pudo completar la lectura. Revisa la conexión del lector o completa los campos manualmente."
        if isinstance(error,ValidationError):
            note = (" · ".join(error.messages)+" Puedes continuar manualmente.")[:350]
    with transaction.atomic():
        changed = DocumentCapture.objects.filter(pk=capture.pk, status="reading", reading_token=token).update(
            status="review", reading_token="", raw_text=text, data=data, service_items=items, extraction=extraction, reading_note=note, version=F("version")+1)
        if changed:
            AuditEvent.objects.create(actor=capture.uploaded_by, entity_type="DocumentCapture", entity_id=str(capture.pk),
                action="read", after={"status": "review", "characters": len(text)})
    return True


def export_content(capture, format):
    """Export text literally so a document cannot inject spreadsheet formulas."""
    data = capture.data
    fields = list(LABELS)
    provenance = {"capture_id": str(capture.pk), "document_type": capture.get_kind_display(),
                  "reviewed_by": capture.reviewed_by.get_full_name() or capture.reviewed_by.username,
                  "reviewed_at": capture.reviewed_at.isoformat()}
    service_fields = ["description", "kind", "quantity", "unit_price", "unit_cost", "part_sku"]
    items = capture.service_items or [{"description":value.strip()} for value in data.get("services", "").splitlines() if value.strip()]
    def literal(value):
        value = "" if value is None else str(value)
        return "'"+value if value.lstrip().startswith(("=", "+", "-", "@")) else value
    if format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        keys = list(provenance) + fields + ["service_" + key for key in service_fields]
        writer.writerow(keys)
        for item in items or [{}]:
            writer.writerow([literal(value) for value in [*provenance.values(),
                *(data.get(key, "") for key in fields), *(item.get(key) for key in service_fields)]])
        return b"\xef\xbb\xbf" + output.getvalue().encode("utf-8"), "text/csv; charset=utf-8"
    book = Workbook()
    order = book.active
    order.title = "Orden revisada"
    order.append(["Campo", "Dato confirmado"])
    for key, value in {**{LABELS[k]:data.get(k, "") for k in fields if k != "services"}, **provenance}.items():
        order.append([key, str(value)])
    services = book.create_sheet("Servicios revisados")
    services.append(["Folio", "Servicio", "Foto origen", "Revisado por", "Tipo",
                     "Cantidad observada", "Precio unitario observado", "Costo unitario observado", "SKU"])
    def observed_number(value):
        if value in (None, ""):
            return None
        try:
            number = Decimal(str(value))
            return number if number.is_finite() else str(value)
        except InvalidOperation:
            return str(value)
    for item in items:
        services.append([data.get("order_number", ""), item.get("description", ""), str(capture.pk),
            provenance["reviewed_by"], item.get("kind"),
            *(observed_number(item.get(key)) for key in ("quantity", "unit_price", "unit_cost")), item.get("part_sku")])
    for row in services.iter_rows(min_row=2):
        row[5].number_format = '#,##0.000'
        for cell in row[6:8]:
            cell.number_format = '#,##0.00'
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 35
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
