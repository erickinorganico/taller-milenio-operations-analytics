import csv
from io import StringIO
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from milenio.client_input import (
    ClientInputError,
    generate_client_templates,
    load_client_input,
    write_client_csv_templates,
    write_client_template,
)


AS_OF = datetime(2026, 1, 15, 12, tzinfo=timezone.utc)


class ClientInputTests(unittest.TestCase):
    def test_sample_workbook_roundtrip_has_canonical_money_and_quality_flags(self):
        with tempfile.TemporaryDirectory() as temp:
            path = write_client_template(Path(temp) / "sample.xlsx", sample=True, as_of=AS_OF, business_name="Taller sintético")
            loaded = load_client_input(path)

        self.assertEqual(loaded["schema_version"], 1)
        self.assertEqual(loaded["metadata"]["as_of"], "2026-01-15T12:00:00Z")
        self.assertEqual(loaded["tables"]["invoices"][0]["amount_cents"], 150000)
        self.assertEqual(loaded["tables"]["payments"][0]["amount_cents"], 50000)
        self.assertEqual(loaded["tables"]["inventory"][0]["unit_cost_cents"], 25000)
        self.assertIn("ORD-004", loaded["metadata"]["quality"]["delivered_orders_without_issued_invoice"])
        self.assertEqual(loaded["metadata"]["quality"]["overdue_partial_invoice_ids"], ["INV-001"])
        self.assertEqual(loaded["metadata"]["quality"]["inventory_shortage_parts"], ["PART-001"])
        self.assertEqual(loaded["metadata"]["quality"]["cancelled_orders_excluded_from_open_work"], 1)

    def test_csv_roundtrip_preserves_optional_blank_as_unknown(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = write_client_csv_templates(Path(temp), sample=True, as_of=AS_OF)
            orders = directory / "orders.csv"
            text = orders.read_text(encoding="utf-8")
            orders.write_text(text.replace("Filtro pendiente,TECH-01", "Filtro pendiente,", 1), encoding="utf-8")
            loaded = load_client_input(directory)

        first = loaded["tables"]["orders"][0]
        self.assertIsNone(first["technician_ref"])
        self.assertIsNone(first["delivered_at"])
        self.assertNotEqual(first["estimated_hours"], 0)

    def test_invalid_money_rejects_nan_and_fractional_cent(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = write_client_csv_templates(Path(temp), sample=True, as_of=AS_OF)
            invoice = directory / "invoices.csv"
            rows = list(csv.reader(invoice.read_text(encoding="utf-8").splitlines()))
            rows[1][5] = "NaN"
            invoice.write_text("\n".join(",".join(row) for row in rows) + "\n", encoding="utf-8")
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(directory)
        self.assertTrue(any(issue["code"] == "invalid_money" for issue in caught.exception.issues))

        with tempfile.TemporaryDirectory() as temp:
            directory = write_client_csv_templates(Path(temp), sample=True, as_of=AS_OF)
            invoice = directory / "invoices.csv"
            rows = list(csv.reader(invoice.read_text(encoding="utf-8").splitlines()))
            rows[1][5] = "1500.001"
            invoice.write_text("\n".join(",".join(row) for row in rows) + "\n", encoding="utf-8")
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(directory)
        self.assertTrue(any(issue["code"] == "fractional_cent" for issue in caught.exception.issues))

    def test_duplicate_and_bad_reference_are_structured(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = write_client_csv_templates(Path(temp), sample=True, as_of=AS_OF)
            invoice = directory / "invoices.csv"
            lines = invoice.read_text(encoding="utf-8").splitlines()
            lines.append(lines[1].replace("INV-001", "INV-002").replace("ORD-002", "ORD-NOPE"))
            invoice.write_text("\n".join(lines) + "\n", encoding="utf-8")
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(directory)
        codes = {issue["code"] for issue in caught.exception.issues}
        self.assertIn("duplicate_ref", codes)
        self.assertIn("bad_ref", codes)

    def test_missing_column_and_cutoff_date_fail_with_cell_context(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = write_client_csv_templates(Path(temp), sample=True, as_of=AS_OF)
            orders = directory / "orders.csv"
            lines = orders.read_text(encoding="utf-8").splitlines()
            lines[0] = lines[0].replace(",Estado", "")
            orders.write_text("\n".join(lines) + "\n", encoding="utf-8")
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(directory)
        self.assertTrue(any(issue["code"] == "missing_column" and issue["sheet"] == "Ordenes" for issue in caught.exception.issues))

        with tempfile.TemporaryDirectory() as temp:
            directory = write_client_csv_templates(Path(temp), sample=True, as_of=AS_OF)
            orders = directory / "orders.csv"
            lines = orders.read_text(encoding="utf-8").splitlines()
            lines[1] = lines[1].replace("2026-01-03T12:00:00Z", "2026-01-16T12:00:00Z")
            orders.write_text("\n".join(lines) + "\n", encoding="utf-8")
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(directory)
        self.assertTrue(any(issue["code"] == "after_as_of" for issue in caught.exception.issues))

    def test_explicit_as_of_override_requires_timezone(self):
        with tempfile.TemporaryDirectory() as temp:
            path = write_client_template(Path(temp) / "sample.xlsx", sample=True, as_of=AS_OF)
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(path, as_of="2026-01-15T12:00:00")
        self.assertTrue(any(issue["code"] == "timezone_required" for issue in caught.exception.issues))

    def test_due_date_is_optional_but_business_name_is_required(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = write_client_csv_templates(Path(temp), sample=True, as_of=AS_OF)
            invoice = directory / "invoices.csv"
            rows = list(csv.reader(invoice.read_text(encoding="utf-8").splitlines()))
            due_index = rows[0].index("Vencimiento")
            for row in rows:
                row.pop(due_index)
            invoice.write_text("\n".join(",".join(row) for row in rows) + "\n", encoding="utf-8")
            loaded = load_client_input(directory)
            self.assertIsNone(loaded["tables"]["invoices"][0]["due_at"])

            config = directory / "config.csv"
            config_rows = list(csv.reader(config.read_text(encoding="utf-8").splitlines()))
            config_rows[1][1] = ""
            config.write_text("\n".join(",".join(row) for row in config_rows) + "\n", encoding="utf-8")
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(directory)
        self.assertTrue(any(issue["code"] == "missing_required" and issue["column"] == "business_name" for issue in caught.exception.issues))

    def test_payment_on_void_invoice_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = write_client_csv_templates(Path(temp), sample=True, as_of=AS_OF)
            invoice = directory / "invoices.csv"
            rows = list(csv.reader(invoice.read_text(encoding="utf-8").splitlines()))
            rows[1][-1] = "void"
            invoice.write_text("\n".join(",".join(row) for row in rows) + "\n", encoding="utf-8")
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(directory)
        self.assertTrue(any(issue["code"] == "payment_on_void" for issue in caught.exception.issues))

    def test_malformed_thousands_and_extreme_decimal_fail_without_coercion(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = write_client_csv_templates(Path(temp), sample=True, as_of=AS_OF)
            invoice = directory / "invoices.csv"
            rows = list(csv.reader(invoice.read_text(encoding="utf-8").splitlines()))
            rows[1][5] = "1,2.34"
            buffer = StringIO(); csv.writer(buffer).writerows(rows)
            invoice.write_text(buffer.getvalue(), encoding="utf-8")
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(directory)
        self.assertTrue(any(issue["code"] == "invalid_money" for issue in caught.exception.issues))

        with tempfile.TemporaryDirectory() as temp:
            directory = write_client_csv_templates(Path(temp), sample=True, as_of=AS_OF)
            invoice = directory / "invoices.csv"
            rows = list(csv.reader(invoice.read_text(encoding="utf-8").splitlines()))
            rows[1][5] = "1e999999"
            buffer = StringIO(); csv.writer(buffer).writerows(rows)
            invoice.write_text(buffer.getvalue(), encoding="utf-8")
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(directory)
        self.assertTrue(any(issue["code"] == "number_range" for issue in caught.exception.issues))

    def test_as_of_offset_normalizes_and_sample_cannot_be_marked_real(self):
        with tempfile.TemporaryDirectory() as temp:
            path = write_client_template(Path(temp) / "sample.xlsx", sample=True, as_of=AS_OF)
            loaded = load_client_input(path, as_of="2026-01-15T18:00:00+06:00")
            self.assertEqual(loaded["metadata"]["as_of"], "2026-01-15T12:00:00Z")
            with self.assertRaises(ValueError):
                write_client_template(Path(temp) / "real-sample.xlsx", sample=True, as_of=AS_OF, synthetic=False)
            with self.assertRaises(ValueError):
                write_client_csv_templates(Path(temp) / "real-sample-csv", sample=True, as_of=AS_OF, synthetic=False)

    def test_blank_template_prefills_business_name_and_duplicate_csv_table_fails(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workbook = write_client_template(root / "blank.xlsx", business_name="Taller cliente")
            book = load_workbook(workbook, read_only=True, data_only=True)
            config_values = list(book["Config"].values)
            book.close()
            self.assertEqual(config_values[3][1], "Taller cliente")
            blank_csv = write_client_csv_templates(root / "blank_csv", business_name="Taller cliente")
            blank_config = list(csv.reader((blank_csv / "config.csv").read_text(encoding="utf-8").splitlines()))
            self.assertEqual(blank_config[1][1], "Taller cliente")

            directory = write_client_csv_templates(root / "sample", sample=True, as_of=AS_OF, business_name="Taller cliente")
            (directory / "Ordenes.csv").write_bytes((directory / "orders.csv").read_bytes())
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(directory)
        self.assertTrue(any(issue["code"] == "duplicate_table" and issue["sheet"] in {"orders.csv", "Ordenes.csv"} for issue in caught.exception.issues))

    def test_duplicate_workbook_alias_fails_without_overwrite(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workbook = write_client_template(root / "sample.xlsx", sample=True, as_of=AS_OF, business_name="Taller cliente")
            book = load_workbook(workbook)
            duplicate = book.create_sheet("Orders")
            duplicate.append(["order_id", "customer_ref", "segment", "received_at", "status"])
            duplicate.append(["ORD-DUP", "CUST-DUP", "particular", "2026-01-01T12:00:00Z", "received"])
            book.save(workbook)
            book.close()
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(workbook)
        self.assertTrue(any(issue["code"] == "duplicate_table" and issue["sheet"] == "Orders" for issue in caught.exception.issues))

    def test_empty_blank_template_fails_closed_instead_of_reporting_zeros(self):
        with tempfile.TemporaryDirectory() as temp:
            path = write_client_template(Path(temp) / "blank.xlsx", sample=False)
            with self.assertRaises(ClientInputError) as caught:
                load_client_input(path)
        self.assertTrue(any(issue["code"] == "no_data" for issue in caught.exception.issues))

    def test_generate_client_templates_produces_two_versions(self):
        with tempfile.TemporaryDirectory() as temp:
            outputs = generate_client_templates(Path(temp), as_of=AS_OF)
            self.assertTrue(outputs["blank_workbook"].exists())
            self.assertTrue(outputs["sample_workbook"].exists())
            self.assertTrue((outputs["sample_csv"] / "orders.csv").exists())
            self.assertEqual(load_client_input(outputs["sample_csv"])["metadata"]["snapshot_id"], "sample-client-001")


if __name__ == "__main__":
    unittest.main()
