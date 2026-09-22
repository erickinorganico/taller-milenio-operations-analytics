import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from milenio.client_actions import (
    ClientActionError,
    WORKBOOK_HEADER_ROW,
    WORKBOOK_HEADERS,
    append_review_log,
    load_action_reviews,
    write_action_workbook,
)


class ClientActionTests(unittest.TestCase):
    def setUp(self):
        self.metadata = {"snapshot_id": "snap-2026-09-21", "as_of": "2026-09-21T18:00:00Z", "synthetic": True}
        self.decisions = [
            {
                "action_id": "ACT-001",
                "category": "collections",
                "title": "Conciliar saldo abierto",
                "priority": "high",
                "owner_role": "administración",
                "why_now": "El corte muestra un saldo pendiente.",
                "recommended_next_step": "Revisar comprobante y registrar la conclusión.",
                "amount_at_risk_cents": 12500,
                "evidence": [{"entity_type": "invoices", "entity_id": "INV-001", "field": "balance_cents", "value": 12500, "version": 1}],
                "source_ids": ["INV-001", "PAY-001"],
            },
            {
                "action_id": "ACT-002",
                "category": "operations",
                "title": "Revisar orden detenida",
                "priority": "medium",
                "owner_role": "jefe de taller",
                "why_now": "La orden conserva un bloqueo al corte.",
                "recommended_next_step": "Confirmar disponibilidad y documentar el siguiente paso.",
                "amount_at_risk_cents": None,
                "evidence": [{"entity_type": "work_orders", "entity_id": "WO-001", "field": "status", "value": "waiting_parts", "version": 1}],
                "source_ids": ["WO-001"],
            },
        ]

    def _write(self, root: Path) -> Path:
        path = root / "seguimiento.xlsx"
        write_action_workbook(path, self.decisions, self.metadata)
        return path

    @staticmethod
    def _columns(sheet):
        return {cell.value: cell.column for cell in sheet[1]}

    def test_edit_excel_row_then_import_valid(self):
        with tempfile.TemporaryDirectory() as temp:
            path = self._write(Path(temp))
            book = openpyxl.load_workbook(path)
            sheet = book["Seguimiento"]
            columns = self._columns(sheet)
            sheet.cell(2, columns[WORKBOOK_HEADERS["owner"]]).value = "Administración"
            sheet.cell(2, columns[WORKBOOK_HEADERS["status"]]).value = "done"
            sheet.cell(2, columns[WORKBOOK_HEADERS["target_date"]]).value = "2026-09-24"
            sheet.cell(2, columns[WORKBOOK_HEADERS["note"]]).value = "Se revisó el saldo en el corte interno."
            sheet.cell(2, columns[WORKBOOK_HEADERS["outcome_evidence"]]).value = "Conciliación INV-001-2026-09"
            book.save(path)

            reviews = load_action_reviews(path, self.decisions, self.metadata)
            self.assertEqual(reviews[0]["action_id"], "ACT-001")
            self.assertEqual(reviews[0]["status"], "done")
            self.assertTrue(reviews[0]["self_reported"])
            self.assertFalse(reviews[0]["external_business_action"])
            self.assertEqual(reviews[1]["status"], "pending")

    def test_forged_source_snapshot_or_title_is_rejected(self):
        for field, value in (("source_snapshot", "snap-forged"), ("title", "Título inventado")):
            with self.subTest(field=field), tempfile.TemporaryDirectory() as temp:
                path = self._write(Path(temp))
                book = openpyxl.load_workbook(path)
                sheet = book["Seguimiento"]
                columns = self._columns(sheet)
                sheet.cell(2, columns[WORKBOOK_HEADERS[field]]).value = value
                book.save(path)
                with self.assertRaises(ClientActionError):
                    load_action_reviews(path, self.decisions, self.metadata)

    def test_unknown_action_id_and_done_without_proof_are_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            path = self._write(Path(temp))
            book = openpyxl.load_workbook(path)
            sheet = book["Seguimiento"]
            columns = self._columns(sheet)
            sheet.cell(2, columns[WORKBOOK_HEADERS["action_id"]]).value = "ACT-UNKNOWN"
            book.save(path)
            with self.assertRaises(ClientActionError):
                load_action_reviews(path, self.decisions, self.metadata)

        with tempfile.TemporaryDirectory() as temp:
            path = self._write(Path(temp))
            book = openpyxl.load_workbook(path)
            sheet = book["Seguimiento"]
            columns = self._columns(sheet)
            sheet.cell(2, columns[WORKBOOK_HEADERS["status"]]).value = "done"
            book.save(path)
            with self.assertRaises(ClientActionError):
                load_action_reviews(path, self.decisions, self.metadata)

    def test_duplicate_import_hash_is_idempotent(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            path = self._write(root)
            book = openpyxl.load_workbook(path)
            sheet = book["Seguimiento"]
            columns = self._columns(sheet)
            sheet.cell(2, columns[WORKBOOK_HEADERS["owner"]]).value = "Administración"
            sheet.cell(2, columns[WORKBOOK_HEADERS["status"]]).value = "accepted"
            sheet.cell(2, columns[WORKBOOK_HEADERS["target_date"]]).value = "2026-09-24"
            book.save(path)
            reviews = load_action_reviews(path, self.decisions, self.metadata)
            log = root / "reviews.jsonl"
            first = append_review_log(log, reviews, "sha256:abc", reviewer="sample-reviewer")
            before = log.read_text(encoding="utf-8")
            second = append_review_log(log, reviews, "sha256:abc", reviewer="sample-reviewer")
            after = log.read_text(encoding="utf-8")
            self.assertEqual(first["status"], "appended")
            self.assertEqual(second["status"], "duplicate")
            self.assertEqual(before, after)
            self.assertEqual(len(after.splitlines()), 1)
            event = json.loads(after)
            self.assertEqual(event["reviewer_claim"], "self_reported")
            self.assertFalse(event["external_business_action"])

    def test_workbook_has_only_expected_sheets_and_headers(self):
        with tempfile.TemporaryDirectory() as temp:
            path = self._write(Path(temp))
            book = openpyxl.load_workbook(path, read_only=True)
            self.assertEqual(book.sheetnames, ["Seguimiento", "Instrucciones"])
            self.assertEqual(tuple(cell.value for cell in book["Seguimiento"][1]), WORKBOOK_HEADER_ROW)
            book.close()

    def test_display_columns_are_manager_friendly_and_lineage_is_hidden(self):
        with tempfile.TemporaryDirectory() as temp:
            path = self._write(Path(temp))
            book = openpyxl.load_workbook(path, read_only=False)
            sheet = book["Seguimiento"]
            self.assertEqual(sheet.freeze_panes, "D2")
            self.assertEqual(sheet["H2"].value, 125.0)
            for column in ("A", "B", "E", "F", "G", "H", "I", "J", "K"):
                self.assertTrue(sheet.column_dimensions[column].hidden)
            for column in ("C", "D", "L", "M", "N", "O", "P"):
                self.assertFalse(sheet.column_dimensions[column].hidden)
            self.assertEqual(sheet.row_dimensions[2].height, 56.0)
            book.close()

    def test_status_specific_human_fields_are_required(self):
        for status, missing_field in (("accepted", "target_date"), ("in_progress", "owner"), ("dismissed", "note")):
            with self.subTest(status=status), tempfile.TemporaryDirectory() as temp:
                path = self._write(Path(temp))
                book = openpyxl.load_workbook(path)
                sheet = book["Seguimiento"]
                columns = self._columns(sheet)
                sheet.cell(2, columns[WORKBOOK_HEADERS["status"]]).value = status
                if status == "accepted":
                    sheet.cell(2, columns[WORKBOOK_HEADERS["owner"]]).value = "Administración"
                elif status == "in_progress":
                    sheet.cell(2, columns[WORKBOOK_HEADERS["target_date"]]).value = "2026-09-24"
                else:
                    sheet.cell(2, columns[WORKBOOK_HEADERS["owner"]]).value = "Administración"
                book.save(path)
                with self.assertRaises(ClientActionError) as caught:
                    load_action_reviews(path, self.decisions, self.metadata)
                self.assertIn(missing_field, str(caught.exception))

    def test_formula_in_annotation_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            path = self._write(Path(temp))
            book = openpyxl.load_workbook(path)
            sheet = book["Seguimiento"]
            columns = self._columns(sheet)
            sheet.cell(2, columns[WORKBOOK_HEADERS["note"]]).value = "=1+1"
            book.save(path)
            with self.assertRaises(ClientActionError):
                load_action_reviews(path, self.decisions, self.metadata)


if __name__ == "__main__":
    unittest.main()
