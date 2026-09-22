import copy
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

import openpyxl

from milenio.client_actions import WORKBOOK_COLUMNS
from milenio.client_decisions import analyze_client, compare_clients
from milenio.client_delivery import build_client, build_comparison, import_client_review, load_client_report, verify_client
from milenio.client_input import load_client_input, write_client_template
from milenio.__main__ import main


class ClientDeliveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'Entrada.xlsx'
        write_client_template(self.source, sample=True, as_of='2026-09-22T18:00:00Z', business_name='Cliente ficticio', snapshot_id='semana-01')
        self.snapshot = load_client_input(self.source)

    def test_money_cancellation_shortage_and_linked_payment_evidence(self):
        a = analyze_client(self.snapshot)
        self.assertEqual(a['summary']['invoiced_cents'], 230000)
        self.assertEqual(a['summary']['paid_cents'], 130000)
        self.assertEqual(a['summary']['balance_cents'], 100000)
        self.assertEqual(a['summary']['overdue_balance_cents'], 100000)
        self.assertEqual(a['summary']['open_orders'], 3)
        collection = next(d for d in a['decisions'] if d['category'] == 'cartera')
        self.assertEqual(collection['amount_at_risk_cents'], 100000)
        self.assertEqual({e['record_id'] for e in collection['evidence'] if e['table'] == 'payments'}, {'PAY-001'})
        self.assertFalse(any('ORD-005' in d['source_ids'] for d in a['decisions']))
        self.assertEqual(next(d for d in a['decisions'] if d['category'] == 'inventario')['priority'], 'P1')

    def test_unknown_not_zero_and_undefined_due_not_overdue(self):
        self.snapshot['tables']['inventory'] = []
        self.snapshot['tables']['invoices'][0]['due_at'] = None
        a = analyze_client(self.snapshot)
        self.assertIsNone(a['summary']['low_stock_parts'])
        self.assertEqual(a['summary']['overdue_balance_cents'], 0)
        self.assertEqual(next(d for d in a['decisions'] if d['category'] == 'vencimiento')['priority'], 'P3')

    def test_comparison_missing_invoice_never_resolved(self):
        before = analyze_client(self.snapshot)
        current = copy.deepcopy(self.snapshot)
        current['metadata'].update(snapshot_id='semana-02',as_of='2026-09-29T18:00:00Z')
        current['tables']['invoices'] = current['tables']['invoices'][1:]
        current['tables']['payments'] = current['tables']['payments'][1:]
        result = compare_clients(before,analyze_client(current))
        item = next(c for c in result['changes'] if c['record_id'] == 'INV-001')
        self.assertEqual(item['state'],'missing_in_later_extract')
        self.assertIsNone(item['observed_payment_increase_cents'])

    def test_comparison_rejects_different_customer_and_clock(self):
        before = analyze_client(self.snapshot)
        with self.assertRaises(ValueError):
            compare_clients(before,before)
        after = copy.deepcopy(before)
        after['metadata'].update(snapshot_id='otro',as_of='2026-09-29T18:00:00Z',business_name='Otro negocio')
        with self.assertRaises(ValueError):
            compare_clients(before,after)

    def test_private_output_boundary(self):
        self.snapshot['metadata']['synthetic'] = False
        with self.assertRaisesRegex(ValueError,'private'):
            build_client(self.snapshot,self.root / 'public',private_root=self.root / 'private')
        result = build_client(self.snapshot,self.root / 'private' / 'corte',private_root=self.root / 'private')
        self.assertFalse(result['synthetic'])

    def test_sealed_files_and_mutable_review(self):
        out = self.root / 'report'
        build_client(self.snapshot,out)
        receipt = verify_client(out)
        self.assertEqual(receipt['snapshot_id'],'semana-01')
        self.assertEqual(receipt['source_hashes'],self.snapshot['source_hashes'])
        (out / 'Seguimiento.xlsx').write_bytes(b'user editable')
        verify_client(out)
        (out / 'analysis.json').write_text('{}',encoding='utf-8')
        with self.assertRaises(ValueError):
            verify_client(out)

    def test_nested_file_and_receipt_identity_tampering_rejected(self):
        out = self.root / 'report'
        build_client(self.snapshot,out)
        (out / 'nested').mkdir()
        added = out / 'nested' / 'extra.txt'
        added.write_text('unexpected',encoding='utf-8')
        with self.assertRaises(ValueError):
            verify_client(out)
        added.unlink()
        receipt = json.loads((out / 'receipt.json').read_text(encoding='utf-8'))
        receipt['synthetic'] = False
        (out / 'receipt.json').write_text(json.dumps(receipt),encoding='utf-8')
        with self.assertRaises(ValueError):
            verify_client(out)

    def test_complete_weekly_journey_with_declared_followup(self):
        out = self.root / 'week1'
        build_client(self.snapshot,out)
        wb = openpyxl.load_workbook(out / 'Seguimiento.xlsx')
        ws = wb['Seguimiento']
        for key,value in [('owner','Responsable ficticio'),('status','in_progress'),('target_date','2026-09-25'),('note','Conciliar el pago con la fuente')]:
            ws.cell(2,WORKBOOK_COLUMNS.index(key)+1,value)
        wb.save(out / 'Seguimiento.xlsx')
        wb.close()
        result = import_client_review(out / 'Seguimiento.xlsx',out,'Revisor ficticio')
        self.assertEqual(result['status'],'appended')
        earlier_bytes = (out / 'Seguimiento.xlsx').read_bytes()
        duplicate = import_client_review(out / 'Seguimiento.xlsx',out,'Revisor ficticio')
        self.assertEqual(duplicate['status'],'duplicate')
        wb = openpyxl.load_workbook(out / 'Seguimiento.xlsx')
        ws = wb['Seguimiento']
        ws.cell(2,WORKBOOK_COLUMNS.index('status')+1,'done')
        ws.cell(2,WORKBOOK_COLUMNS.index('outcome_evidence')+1,'Referencia ficticia de verificación')
        wb.save(out / 'Seguimiento.xlsx'); wb.close()
        import_client_review(out / 'Seguimiento.xlsx',out,'Segundo revisor ficticio')
        old_book = self.root / 'old.xlsx'
        old_book.write_bytes(earlier_bytes)
        import_client_review(old_book,out,'Etiqueta antigua')
        self.assertIn('Segundo revisor ficticio',(out / 'REVISION.html').read_text(encoding='utf-8'))
        self.assertNotIn('Etiqueta antigua',(out / 'REVISION.html').read_text(encoding='utf-8'))
        current = copy.deepcopy(self.snapshot)
        current['metadata'].update(snapshot_id='semana-02',as_of='2026-09-29T18:00:00Z')
        current['tables']['payments'].append({'payment_id':'PAY-003','invoice_id':'INV-001','paid_at':'2026-09-24T12:00:00Z','amount_cents':40000,'method':'transfer'})
        later = self.root / 'week2'
        build_client(current,later)
        build_comparison(out,later,self.root / 'comparison')
        comparison = json.loads((self.root / 'comparison' / 'comparison.json').read_text(encoding='utf-8'))
        invoice = next(c for c in comparison['changes'] if c['record_id'] == 'INV-001')
        self.assertEqual(invoice['after']['balance_cents'],60000)
        self.assertEqual(invoice['observed_payment_increase_cents'],40000)
        self.assertTrue(any(c['human_review'] and c['human_review']['owner'] == 'Responsable ficticio' for c in comparison['action_continuity']))
        verify_client(out)

    def test_html_escapes_source_and_excel_has_no_formulas(self):
        self.snapshot['metadata']['business_name'] = '<script>alert(1)</script>'
        self.snapshot['tables']['inventory'][0]['description'] = '=HYPERLINK("https://example.invalid","x")'
        out = self.root / 'report'
        build_client(self.snapshot,out)
        self.assertNotIn('<script>alert(1)</script>',(out / 'INICIO.html').read_text(encoding='utf-8'))
        wb = openpyxl.load_workbook(out / 'Gerencia.xlsx',data_only=False)
        self.assertEqual(len(wb.sheetnames),7)
        self.assertFalse(any(c.data_type == 'f' for ws in wb for row in ws for c in row))
        wb.close()

    def test_cli_invalid_input_produces_no_false_report(self):
        blank = self.root / 'blank.xlsx'
        write_client_template(blank)
        stream = io.StringIO()
        with redirect_stdout(stream):
            code = main(['client-analyze','--input',str(blank),'--output',str(self.root / 'invalid')])
        self.assertEqual(code,2)
        self.assertFalse((self.root / 'invalid').exists())
        self.assertFalse(json.loads(stream.getvalue())['report_generated'])


if __name__ == '__main__':
    unittest.main()
