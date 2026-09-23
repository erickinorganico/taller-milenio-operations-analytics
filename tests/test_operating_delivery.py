import hashlib
import json
import shutil
import tempfile
import unittest
from pathlib import Path

import openpyxl

from milenio.client_actions import WORKBOOK_HEADERS
from milenio.operating_delivery import import_operating_review
from milenio.studio import build_studio, verify_studio


class OperatingDeliveryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name)
        cls.output = cls.root / 'studio'
        cls.result = build_studio(cls.output)
        cls.model = json.loads((cls.output / 'operating_model.json').read_text(encoding='utf8'))

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_all_source_rows_metric_ids_and_decisions_have_resolvable_lineage(self):
        data = self.model
        tables = {t['name']:t for t in data['sources']['tables']}
        metrics = {m['metric_id']:m for m in data['registry']['metrics']}
        self.assertEqual((len(tables),len(metrics),len(data['agents']['agents'])),(33,31,9))
        digest = hashlib.sha256((self.output / 'warehouse.sqlite').read_bytes()).hexdigest()
        self.assertEqual(data['metadata']['snapshot_id'],digest)
        self.assertEqual(data['registry']['source_sha256'],digest)
        self.assertEqual(sum(t['row_count'] for t in tables.values() if t['kind']=='physical_source'),self.result['source_rows'])
        for metric in metrics.values():
            for name in metric['source_tables']:
                self.assertIn(metric['metric_id'],tables[name]['metric_ids'])
        for decision in data['decisions']:
            self.assertTrue(decision['evidence'])
            for mid in decision['metric_ids']: self.assertIn(mid,metrics)
            for ref in decision['evidence']:
                if 'table' not in ref: continue  # separately typed process contract references
                source = tables[ref['table']]
                row = next(r for r in source['rows'] if str(r[source['logical_key'][0]])==str(ref['record_id']))
                self.assertEqual(row[ref['field']],ref['value'])
        html = (self.output / 'INICIO.html').read_text(encoding='utf8')
        self.assertIn('application/json',html)
        self.assertNotIn('__PAYLOAD__',html)
        self.assertNotIn(str(self.root),html)
        self.assertEqual(verify_studio(self.output)['status'],'pass')

    def test_human_review_import_preserves_source_and_rejects_forged_origin(self):
        bookpath = self.root / 'review-copy.xlsx'
        shutil.copyfile(self.output / 'Seguimiento.xlsx', bookpath)
        book = openpyxl.load_workbook(bookpath)
        try:
            sheet=book['Seguimiento']; cols={c.value:c.column for c in sheet[1]}
            for key,value in {'owner':'Revisor de prueba','status':'accepted','target_date':'2026-09-24','note':'Verificar el caso con fuente.'}.items():
                sheet.cell(2,cols[WORKBOOK_HEADERS[key]]).value=value
            book.save(bookpath)
        finally: book.close()
        before=(self.output / 'receipt.json').read_bytes()
        result=import_operating_review(self.output,bookpath,self.root/'human-review','Revisor de prueba')
        self.assertFalse(result['external_business_action'])
        self.assertEqual(result['reviews'],len(self.model['decisions']))
        self.assertEqual(before,(self.output / 'receipt.json').read_bytes())
        self.assertEqual(verify_studio(self.output)['status'],'pass')
        with self.assertRaises(ValueError):
            import_operating_review(self.output,bookpath,self.output/'invalid-review','Revisor')
        book=openpyxl.load_workbook(bookpath)
        try:
            book['Seguimiento'].cell(2,cols[WORKBOOK_HEADERS['title']]).value='Origen alterado'
            book.save(bookpath)
        finally: book.close()
        with self.assertRaises(ValueError):
            import_operating_review(self.output,bookpath,self.root/'forged-review','Revisor')
        self.assertFalse((self.root/'forged-review').exists())


if __name__=='__main__': unittest.main()
